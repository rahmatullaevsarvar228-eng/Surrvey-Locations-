# -*- coding: utf-8 -*-
"""Excel-отчёты в едином оформлении (как в main.py): тёмно-синяя шапка,
заливка строк по статусу 🔴/🟡/🟢, закреплённая шапка, автофильтр, ширина
колонок по содержимому."""
import io
import re

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
FILLS = {
    "RED": PatternFill("solid", fgColor="FCECEA"),
    "YELLOW": PatternFill("solid", fgColor="FFF8E6"),
    "GREEN": PatternFill("solid", fgColor="EAF5F0"),
}
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE_FONT = Font(bold=True, size=12, color="1F4E78")
STATUS_TEXT = {"RED": "🔴 RED", "YELLOW": "🟡 YELLOW", "GREEN": "🟢 GREEN"}


def _sheet_name(name):
    return re.sub(r"[\\/*?:\[\]]", "", name)[:31] or "Лист"


def write_sheet(writer, name, df, status_col=None, row_status=None, title=None):
    """Пишет таблицу и оформляет её. row_status — список кодов RED/YELLOW/GREEN
    по строкам (если заливка зависит не от видимой колонки)."""
    name = _sheet_name(name)
    df = df.copy()
    if status_col and status_col in df.columns:
        row_status = row_status or df[status_col].tolist()
        df[status_col] = df[status_col].map(lambda s: STATUS_TEXT.get(s, s))
    start = 2 if title else 0
    df.to_excel(writer, index=False, sheet_name=name, startrow=start)
    ws = writer.sheets[name]
    if title:
        ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    header_row = start + 1
    ncol = len(df.columns)
    for c in range(1, ncol + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in range(len(df)):
        fill = FILLS.get(row_status[r]) if row_status else None
        for c in range(1, ncol + 1):
            cell = ws.cell(row=header_row + 1 + r, column=c)
            cell.border = BORDER
            if fill:
                cell.fill = fill
    for i, col in enumerate(df.columns, start=1):
        longest = max([len(str(col))] + [len(str(v)) for v in df[col].head(500)])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, longest + 2))
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if len(df):
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncol)}{header_row + len(df)}"
    return ws


def _to_bytes(build):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        build(writer)
    return buf.getvalue()


def defects_table(df, only_defects=True):
    rows = df[df["is_defect"]] if only_defects else df
    out = pd.DataFrame({
        "ID анкеты": rows["row_id"], "Город": rows["city"], "Интервьюер": rows["inter"],
        "Device ID": rows["deviceid"],
        "Старт": rows["start"].dt.strftime("%d.%m.%Y %H:%M"),
        "Финиш": rows["end"].dt.strftime("%d.%m.%Y %H:%M"),
        "Длительность, мин": rows["duration_min"].round(1),
        "Брак?": rows["is_defect"].map({True: "Да", False: "Нет"}),
        "Причина брака": rows["reason_text"], "Предупреждения": rows["warning_text"],
    })
    return out.sort_values(["Город", "Интервьюер"], key=lambda s: s.astype(str)).reset_index(drop=True)


def full_report(result, legend):
    """Полный отчёт по волне: брак, все анкеты, интервьюеры, города, повтор."""
    df = result["df"]

    def build(writer):
        defects = defects_table(df)
        write_sheet(writer, "Брак", defects, row_status=["RED"] * len(defects))
        inter = pd.DataFrame(result["interviewers"])
        if len(inter):
            write_sheet(writer, "Интервьюеры", inter.drop(columns=["Повтор: статус"]), status_col="Статус")
        leg = pd.DataFrame([{"Статус": x["code"], "Описание": x["desc"], "Действия": x["actions"]} for x in legend])
        write_sheet(writer, "Легенда статусов", leg, status_col="Статус")
        if result["city_issues"]:
            ci = pd.DataFrame([{"Город": x["city"], "Уровень": x["level"], "Проблема": x["text"]}
                               for x in result["city_issues"]])
            write_sheet(writer, "Города", ci, status_col="Уровень")
        if result["repetition"]["enabled"] and result["repetition"]["rows"]:
            write_sheet(writer, "Повтор значения", pd.DataFrame(result["repetition"]["rows"]),
                        status_col="Статус")
        g = result.get("geo") or {}
        if g.get("enabled") and g.get("points"):
            pts = pd.DataFrame(g["points"])
            gps = pd.DataFrame({
                "ID анкеты": pts["id"], "Город": pts["city"], "Интервьюер": pts["inter"],
                "Широта": pts["lat"], "Долгота": pts["lon"], "Ближайшая точка": pts["point"],
                "Расстояние, км": pts["dist"],
                "Далеко от точки": pts["far"].map({True: "Да", False: ""}),
                "В скоплении": pts["cluster"].map({True: "Да", False: ""}),
                "Одинаковые координаты": pts["same"].map({True: "Да", False: ""}),
            })
            status = ["RED" if f or c or s else "GREEN" for f, c, s in zip(pts["far"], pts["cluster"], pts["same"])]
            write_sheet(writer, "GPS", gps, row_status=status)
            if g.get("clusters"):
                cl = pd.DataFrame(g["clusters"]).rename(columns={"lat": "Широта", "lon": "Долгота"})
                write_sheet(writer, "GPS скопления", cl, status_col="Статус")
        all_rows = defects_table(df, only_defects=False)
        status = ["RED" if d == "Да" else ("YELLOW" if w else "GREEN")
                  for d, w in zip(all_rows["Брак?"], all_rows["Предупреждения"])]
        write_sheet(writer, "Все анкеты", all_rows, row_status=status)

    return _to_bytes(build)


def clean_report(clean, todo):
    """Чистая база для заказчика: только принятые и анкеты без замечаний, все
    исходные колонки. Нерешённые подозрительные — отдельным листом."""
    def build(writer):
        write_sheet(writer, "Чистая база", clean)
        if len(todo):
            write_sheet(writer, "Не решено", todo, row_status=["YELLOW"] * len(todo))
    return _to_bytes(build)


def simple_report(sheet, rows, status_col=None):
    df = pd.DataFrame(rows)
    return _to_bytes(lambda w: write_sheet(w, sheet, df, status_col=status_col))


def history_report(hist):
    wave_names = [w["wave"] for w in hist["waves"]]
    rows = []
    for r in hist["rows"]:
        row = {k: v for k, v in r.items() if k != "waves"}
        for w in wave_names:
            x = r["waves"].get(w)
            row[w] = f"{STATUS_TEXT.get(x['status'], x['status'])} · {x['pct']}%" if x else ""
        rows.append(row)
    df = pd.DataFrame(rows)
    status = [r["Худший статус"] if r["Подряд в зоне риска"] else "GREEN" for r in hist["rows"]]
    return _to_bytes(lambda w: write_sheet(w, "История интервьюеров", df, status_col="Худший статус",
                                           row_status=status))
