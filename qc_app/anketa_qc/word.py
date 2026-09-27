# -*- coding: utf-8 -*-
"""Отчёт контроля качества в Word (.docx).

Короткий и понятный документ для руководителя и заказчика: итог одной
фразой, главные цифры, графики, где и почему брак, GPS, интервьюеры, квоты,
выводы и пустое место для комментария руководителя. Всё можно поправить
вручную в Word и сохранить в PDF.

Состояние анкеты считается так же, как в дашбордах: решение руководителя
важнее системы («Принять» убирает из брака, «Брак» добавляет, «На
перезвон» — «проверить»). Технические записи в брак и проценты не входят.
"""
import io
from collections import Counter, defaultdict
from datetime import datetime

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

from . import engine, review  # noqa: E402
from .review import final_state  # noqa: E402,F401

NAVY = RGBColor(0x1F, 0x4E, 0x78)
GREY = RGBColor(0x6E, 0x6E, 0x73)
C_BRAK, C_WARN, C_OK, C_BLUE, C_TECH = "#d03b3b", "#fab219", "#0ca30c", "#2a78d6", "#4a3aa7"
STATUS_TEXT = {"RED": "🔴 Критично", "YELLOW": "🟡 Внимание", "GREEN": "🟢 Норма"}
STATUS_FILL = {"RED": "FDECEB", "YELLOW": "FFF5E0", "GREEN": "E8F6EC"}

plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 9, "axes.spines.top": False, "axes.spines.right": False,
    "axes.edgecolor": "#c3c2b7", "axes.labelcolor": "#52514e", "xtick.color": "#52514e", "ytick.color": "#52514e",
    "axes.grid": True, "grid.color": "#e1e0d9", "grid.linewidth": 0.8, "axes.axisbelow": True,
})


# ─────────────────────────────────────────────────────────────────────────
# Данные
# ─────────────────────────────────────────────────────────────────────────
def frame(result, decisions_by_pos, filters=None):
    """Таблица анкет для отчёта с итоговым состоянием и фильтрами."""
    df = result["df"].copy()
    dec = {p: (d or {}).get("decision") for p, d in decisions_by_pos.items()}
    df["decision"] = df["pos"].map(dec).where(lambda x: x.notna(), None)
    df["state"] = [final_state(t, d, w, x) for t, d, w, x in
                   zip(df["technical"], df["is_defect"], df["is_warning"], df["decision"])]
    for key, col in (("region", "region"), ("city", "city"), ("inter", "inter")):
        v = (filters or {}).get(key)
        if v:
            df = df[df[col] == v]
    return df


def reason_counts(df, sev="defect"):
    c = Counter()
    for xs in df.loc[df["state"] == "brak", "issues"]:
        for code in {code for code, s, _ in xs if s == sev}:
            c[code] += 1
    return c.most_common()


def block_counts(df):
    c = Counter()
    for xs, bl in zip(df.loc[df["state"] == "brak", "issues"], df.loc[df["state"] == "brak", "blocks"]):
        for b in {b for (code, s, _), b in zip(xs, bl) if s == engine.DEFECT}:
            c[b] += 1
    return c.most_common()


def anket(n):
    return f"{n} {engine.plural(n, 'анкета', 'анкеты', 'анкет')}"


def level(pct, cfg):
    st = cfg["status"]
    return "RED" if pct >= st["red_pct"] else "YELLOW" if pct >= st["yellow_pct"] else "GREEN"


# ─────────────────────────────────────────────────────────────────────────
# Оформление
# ─────────────────────────────────────────────────────────────────────────
def _insert_ordered(parent, child, successors):
    """Вставляет элемент перед первым из «последующих» по схеме Word — иначе
    Word считает файл повреждённым."""
    for tag in successors:
        found = parent.find(qn(tag))
        if found is not None:
            found.addprevious(child)
            return
    parent.append(child)


def _shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    _insert_ordered(tc_pr, shd, ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"))


def _borders(table, color="D9D9D9"):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)
    _insert_ordered(tbl_pr, borders, ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption"))


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    _insert_ordered(tr_pr, el, ("w:tblCellSpacing", "w:jc", "w:hidden"))


def add_table(doc, headers, rows, widths=None, status_col=None, font_size=9, align_right=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(t)
    hdr = t.rows[0]
    _repeat_header(hdr)
    for i, name in enumerate(headers):
        cell = hdr.cells[i]
        _shade(cell, "1F4E78")
        p = cell.paragraphs[0]
        r = p.add_run(str(name))
        r.bold, r.font.size = True, Pt(font_size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            p = cells[i].paragraphs[0]
            r = p.add_run("" if v is None else str(v))
            r.font.size = Pt(font_size)
            if i in align_right:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if status_col is not None:
            code = row[status_col]
            for key, fill in STATUS_FILL.items():
                if key in str(code) or STATUS_TEXT[key] == code:
                    for c in cells:
                        _shade(c, fill)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def heading(doc, text, lvl=1):
    h = doc.add_heading(text, level=lvl)
    for r in h.runs:
        r.font.color.rgb = NAVY
    return h


def para(doc, text, size=10.5, color=None, bold=False, italic=False, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold, r.italic = bold, italic
    if color is not None:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    return p


def picture(doc, fig, width_cm=16.5):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    doc.add_picture(buf, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


# ─────────────────────────────────────────────────────────────────────────
# Графики
# ─────────────────────────────────────────────────────────────────────────
def chart_days(df):
    iv = df[df["state"] != "tech"].dropna(subset=["start"])
    if iv.empty:
        return None
    days = iv["start"].dt.normalize()
    tab = pd.crosstab(days, iv["state"]).reindex(columns=["brak", "warn", "ok"], fill_value=0)
    fig, ax = plt.subplots(figsize=(8, 2.8))
    x = np.arange(len(tab))
    bottom = np.zeros(len(tab))
    for col, color, name in (("brak", C_BRAK, "Брак"), ("warn", C_WARN, "Проверить"), ("ok", C_OK, "Норма")):
        ax.bar(x, tab[col], bottom=bottom, color=color, width=0.6, label=name, edgecolor="white", linewidth=1.2)
        bottom += tab[col].to_numpy()
    ax.set_xticks(x, [d.strftime("%d.%m") for d in tab.index], rotation=0 if len(x) < 16 else 90)
    ax.grid(axis="x", visible=False)
    ax.set_ylabel("анкет")
    ax.legend(ncol=3, frameon=False, loc="upper left", bbox_to_anchor=(0, 1.18))
    return fig


def chart_hbars(items, color, value_fmt="{:.0f}", max_value=None, xlabel=""):
    items = items[:12]
    if not items:
        return None
    labels = [str(k) for k, _ in items][::-1]
    vals = [v for _, v in items][::-1]
    fig, ax = plt.subplots(figsize=(8, 0.42 * len(items) + 0.6))
    ax.barh(labels, vals, color=color, height=0.55)
    top = max_value or max(vals) * 1.15 or 1
    ax.set_xlim(0, top)
    for i, v in enumerate(vals):
        ax.text(v + top * 0.01, i, value_fmt.format(v), va="center", fontsize=8.5, color="#333")
    ax.grid(axis="y", visible=False)
    ax.set_xlabel(xlabel)
    return fig


def chart_city_map(city, sub, plan_points, g):
    """Карта города без подложки: плановые точки с радиусами и анкеты."""
    fig, ax = plt.subplots(figsize=(7.5, 5.2))
    lat0 = float(sub["lat"].mean())
    kx = np.cos(np.radians(lat0))
    for p in plan_points:
        r = float(p.get("radius_km") or g.get("max_dist_km", 2.0)) / 111.0
        ax.add_patch(plt.Circle((float(p["lon"]) * kx, float(p["lat"])), r, color=C_BLUE, alpha=0.07, lw=0))
        ax.add_patch(plt.Circle((float(p["lon"]) * kx, float(p["lat"])), r, fill=False, color=C_BLUE, alpha=0.4, lw=0.8))
        ax.scatter(float(p["lon"]) * kx, float(p["lat"]), s=40, marker="^", color=C_BLUE, edgecolor="white", zorder=4)
    for state, color, name, z in (("ok", C_OK, "Норма", 2), ("warn", C_WARN, "Проверить", 3), ("brak", C_BRAK, "Брак", 5),
                                  ("tech", C_TECH, "Тех. записи", 3)):
        s = sub[sub["state"] == state]
        if len(s):
            ax.scatter(s["lon"] * kx, s["lat"], s=18, color=color, edgecolor="white", linewidth=0.5, label=name, zorder=z)
    out = sub[sub["issues"].map(lambda xs: any(c == "geo_city" for c, _, _ in xs))]
    if len(out):
        ax.scatter(out["lon"] * kx, out["lat"], s=90, facecolor="none", edgecolor=C_BRAK, linewidth=1.4,
                   label="Не в своём городе", zorder=6)
    ax.scatter([], [], marker="^", color=C_BLUE, label="Плановая точка")
    ax.set_aspect("equal")
    ax.set_xticks([])
    ax.set_yticks([])
    ax.grid(False)
    for sp in ax.spines.values():
        sp.set_visible(True)
        sp.set_color("#e1e0d9")
    ax.set_title(city, loc="left", fontsize=11, color="#1F4E78", fontweight="bold")
    ax.legend(frameon=False, fontsize=8, loc="upper left", bbox_to_anchor=(1.0, 1.0))
    return fig


# ─────────────────────────────────────────────────────────────────────────
# Документ
# ─────────────────────────────────────────────────────────────────────────
def build(result, cfg, decisions_by_pos, project, source, quotas_result=None, user=None, filters=None):
    df = frame(result, decisions_by_pos, filters)
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:                       # в шаблоне python-docx без обязательного percent
        zoom.set(qn("w:percent"), "100")
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(10.5)

    iv = df[df["state"] != "tech"]
    n_iv, n_tech = len(iv), int((df["state"] == "tech").sum())
    cnt = iv["state"].value_counts()
    n_brak, n_warn, n_ok = int(cnt.get("brak", 0)), int(cnt.get("warn", 0)), int(cnt.get("ok", 0))
    pct = n_brak / n_iv * 100 if n_iv else 0.0
    scope = ", ".join(v for v in (filters or {}).values() if v)
    period = engine.data_period(df) or "—"

    # ── Титул ────────────────────────────────────────────────────────────
    for _ in range(5):
        doc.add_paragraph()
    para(doc, "ОТЧЁТ КОНТРОЛЯ КАЧЕСТВА", 12, GREY, bold=True)
    para(doc, project + (f" — {scope}" if scope else ""), 26, NAVY, bold=True, space_after=12)
    para(doc, f"Период полевых работ: {period}", 12)
    para(doc, f"Анкет проверено: {n_iv}" + (f" (+ {n_tech} технических записей)" if n_tech else ""), 12)
    para(doc, f"Источник: {source or '—'}", 10, GREY)
    para(doc, f"Подготовлено: {datetime.now():%d.%m.%Y %H:%M}" + (f" · {user}" if user else ""), 10, GREY)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── 1. Итог ──────────────────────────────────────────────────────────
    st_cfg = cfg["status"]
    lvl = "RED" if pct >= st_cfg["yellow_pct"] else "YELLOW" if pct >= st_cfg["yellow_pct"] / 2 else "GREEN"
    verdict = {"RED": "Много брака — нужно вмешаться", "YELLOW": "Есть проблемы — стоит проверить",
               "GREEN": "Поле идёт нормально"}[lvl]
    heading(doc, "1. Итог")
    para(doc, f"{STATUS_TEXT[lvl][:2]} {verdict}", 16, bold=True)
    city_stats = []
    for city, g in iv.groupby("city"):
        b = int((g["state"] == "brak").sum())
        city_stats.append((city, g["region"].iloc[0], len(g), b, b / len(g) * 100))
    city_stats.sort(key=lambda x: -x[4])
    reasons = reason_counts(df)
    summary = [f"Брак — {pct:.1f}% ({n_brak} из {n_iv} анкет)."]
    if n_brak and city_stats and city_stats[0][4] > 0:
        summary.append(f"Больше всего — {city_stats[0][0]} ({city_stats[0][4]:.0f}%).")
    if reasons:
        summary.append(f"Главная причина: «{engine.short_label(reasons[0][0])}» ({anket(reasons[0][1])}).")
    para(doc, " ".join(summary), 11)

    add_table(doc, ["Анкет", "Брак", "Проверить", "Норма", "Тех. записи"],
              [[n_iv, f"{n_brak} ({pct:.1f}%)", n_warn, n_ok, n_tech]], align_right=(0, 1, 2, 3, 4), font_size=10)
    dec_counts = Counter(v.get("decision") for v in decisions_by_pos.values() if v and v.get("decision"))
    if dec_counts:
        para(doc, "Решения руководителя проекта: " + ", ".join(f"{k} — {v}" for k, v in dec_counts.items())
             + ". Они уже учтены в цифрах: «Принять» убирает анкету из брака, «Брак» добавляет.", 9.5, GREY)
    fig = chart_days(df)
    if fig is not None:
        para(doc, "Анкеты по дням", 11, NAVY, bold=True, space_after=2)
        picture(doc, fig)

    # ── 2. Где брак ───────────────────────────────────────────────────────
    heading(doc, "2. Где брак: регионы и города")
    reg = defaultdict(lambda: [0, 0])
    for city, region, n, b, _ in city_stats:
        reg[region][0] += n
        reg[region][1] += b
    reg_rows = sorted(((r, n, b, b / n * 100) for r, (n, b) in reg.items()), key=lambda x: -x[3])
    add_table(doc, ["Регион", "Анкет", "Брак", "% брака", "Статус"],
              [[r, n, b, f"{p:.1f}%", STATUS_TEXT[level(p, cfg)]] for r, n, b, p in reg_rows],
              widths=[6, 2.5, 2.5, 2.5, 3.5], status_col=4, align_right=(1, 2, 3))
    fig = chart_hbars([(c, p) for c, _, _, _, p in city_stats], C_BRAK, "{:.0f}%", 100, "% брака среди анкет города")
    if fig is not None:
        picture(doc, fig, 15)
    add_table(doc, ["Город", "Регион", "Анкет", "Брак", "% брака"],
              [[c, r, n, b, f"{p:.1f}%"] for c, r, n, b, p in city_stats], widths=[4, 4.5, 2.5, 2.5, 2.5],
              align_right=(2, 3, 4))

    # ── 3. Почему ─────────────────────────────────────────────────────────
    heading(doc, "3. Почему брак")
    para(doc, "Сколько бракованных анкет с каждой причиной (у одной анкеты может быть несколько причин) "
              "и как система это определяет.", 10, GREY)
    rows = []
    for code, n in reasons:
        logic, _ = review.explain(cfg, code, "", list(result["raw"].columns))
        rows.append([engine.short_label(code), n, f"{n / max(n_brak, 1) * 100:.0f}%", logic])
    if rows:
        fig = chart_hbars([(r[0], r[1]) for r in rows], C_BRAK, "{:.0f}", None, "анкет с браком")
        picture(doc, fig, 15)
        add_table(doc, ["Причина", "Анкет", "% брака", "Как определяется"], rows, widths=[4, 1.6, 1.8, 9.6],
                  align_right=(1, 2), font_size=8.5)
    else:
        para(doc, "Брака нет.")
    blocks = block_counts(df)
    if blocks:
        para(doc, "В каких блоках анкеты ошибки", 11, NAVY, bold=True, space_after=2)
        add_table(doc, ["Блок анкеты / область", "Анкет с браком"], [[b, n] for b, n in blocks], widths=[11, 4],
                  align_right=(1,))

    # ── 4. GPS ────────────────────────────────────────────────────────────
    g = result.get("geo") or {}
    if g.get("enabled"):
        heading(doc, "4. GPS: где проводились интервью")
        has = lambda code: df["issues"].map(lambda xs: any(c == code for c, _, _ in xs))  # noqa: E731
        gps_rows = [["С координатами", int(df["lat"].notna().sum())],
                    ["Не в своём городе", int(has("geo_city").sum())],
                    ["Далеко от плановой точки", int(has("geo_far").sum())],
                    ["«Телепорт» (слишком быстрое перемещение)", int(has("geo_jump").sum())],
                    ["Одинаковые координаты", int(has("geo_same").sum())],
                    ["Скопления в одном месте", int(has("geo_cluster").sum())],
                    ["Без координат", int(df["lat"].isna().sum())]]
        add_table(doc, ["Показатель", "Анкет"], gps_rows, widths=[11, 4], align_right=(1,))
        plan = {str(k).lower(): v for k, v in (g.get("plan") or {}).items()}
        geo_issue = df["issues"].map(lambda xs: sum(1 for c, _, _ in xs if c.startswith("geo_")))
        ranked = df[df["lat"].notna()].assign(gi=geo_issue).groupby("city")["gi"].sum().sort_values(ascending=False)
        shown = [c for c in ranked.index if ranked[c] > 0][:4] or list(ranked.index[:2])
        for city in shown:
            sub = df[(df["city"] == city) & df["lat"].notna()]
            pts = (plan.get(str(city).lower()) or {}).get("points") or []
            picture(doc, chart_city_map(city, sub, pts, g), 15.5)
        stats = g.get("point_stats") or []
        if stats:
            para(doc, "Плановые точки: план и факт", 11, NAVY, bold=True, space_after=2)
            add_table(doc, ["Город", "Точка", "План", "Факт", "Интервьюеров"],
                      [[p["Город"], p["Точка"], p.get("Квота") if p.get("Квота") is not None else "—", p["Анкет"],
                        p["Интервьюеров"]] for p in sorted(stats, key=lambda p: (str(p["Город"]), -p["Анкет"]))
                       if not (filters or {}).get("city") or str(p["Город"]).lower() == str(filters["city"]).lower()],
                      widths=[3, 7.5, 2, 2, 2.5], align_right=(2, 3, 4), font_size=8.5)

    # ── 5. Интервьюеры ────────────────────────────────────────────────────
    heading(doc, "5. Интервьюеры")
    inter_rows = []
    for inter, gi in iv.groupby(iv["inter"].fillna("—")):
        b = int((gi["state"] == "brak").sum())
        p = b / len(gi) * 100
        top = ", ".join(f"{engine.short_label(c)} ×{k}" for c, k in reason_counts(gi.assign(state=gi["state"]))[:2])
        inter_rows.append([STATUS_TEXT[level(p, cfg)], inter, gi["city"].mode().iloc[0], len(gi), b, f"{p:.1f}%", top, p])
    inter_rows.sort(key=lambda r: -r[-1])
    sc = Counter(r[0] for r in inter_rows)
    para(doc, "  ·  ".join(f"{STATUS_TEXT[k]}: {sc.get(STATUS_TEXT[k], 0)}" for k in ("RED", "YELLOW", "GREEN")), 11, bold=True)
    add_table(doc, ["Статус", "Интервьюер", "Город", "Анкет", "Брак", "% брака", "Главные причины"],
              [r[:-1] for r in inter_rows], widths=[2.6, 2.4, 2.4, 1.4, 1.4, 1.6, 5.2], status_col=0,
              align_right=(3, 4, 5), font_size=8.5)
    pats = [p for p in result.get("answer_patterns") or []
            if not (filters or {}).get("inter") or p["inter"] == filters["inter"]]
    if pats:
        para(doc, "Необычные ответы респондентов (сравнение с коллегами)", 11, NAVY, bold=True, space_after=2)
        for p in pats[:10]:
            bullet(doc, " " + "; ".join(p["notes"][:2]), f"{p['inter']} ({p['city']}):")
    for code in ("RED", "YELLOW"):
        lv = engine.STATUS_LEVELS[code]
        names = [r[1] for r in inter_rows if r[0] == STATUS_TEXT[code]]
        if names:
            bullet(doc, f" {', '.join(map(str, names))}. Что делать: {lv['actions']}.", f"{STATUS_TEXT[code]}:")

    # ── 6. По дням ────────────────────────────────────────────────────────
    from . import daily
    d_rows, d_days, d_per_day, norm = daily.table(result, cfg, decisions_by_pos, filters)
    n_sec = 6
    if d_days:
        heading(doc, f"{n_sec}. Работа по дням")
        para(doc, (f"Норма — {norm} засчитанных анкет в день (брак в норму не идёт)." if norm else
                   "Дневная норма не задана — показано, сколько засчитанных анкет (без брака) сделано в каждый день."),
             10, GREY)
        add_table(doc, ["День", "Работали", "Выполнили норму", "Ниже нормы", "Не работали", "Засчитано", "Брак"],
                  [[pd.Timestamp(x["day"]).strftime("%d.%m.%Y"), f"{x['worked']} из {x['team']}",
                    x["norm_ok"] if norm else "—", x["low"] if norm else "—", x["off"], x["ok"], x["brak"]]
                   for x in d_per_day], align_right=(2, 3, 4, 5, 6), font_size=9)
        weak = [r for r in d_rows if r["days_low"] or r["days_off"]]
        if weak:
            para(doc, "Кто не выполнял норму или пропускал дни", 11, NAVY, bold=True, space_after=2)
            add_table(doc, ["Интервьюер", "Город", "Работал дней", "Не работал", "Ниже нормы", "В среднем в день"],
                      [[r["inter"], r["city"], r["days_worked"], r["days_off"], r["days_low"] if norm else "—", r["avg"]]
                       for r in sorted(weak, key=lambda r: (-r["days_off"] - r["days_low"], str(r["inter"])))],
                      align_right=(2, 3, 4, 5), font_size=9)
        n_sec += 1

    # ── Квоты ─────────────────────────────────────────────────────────────
    q = quotas_result or {}
    if q.get("enabled"):
        heading(doc, f"{n_sec}. Выполнение квот")
        sm = q["summary"]
        in_plan = sm.get("ok_in_plan", sm["ok"])
        para(doc, f"План выполнен на {sm['pct']}%: в пределах плана засчитано {in_plan} из {sm['plan']}, "
                  f"осталось добрать {sm['left']}"
             + (f", перебор {sm['over']}" if sm.get("over") else "") + ".", 11)
        labels = q["labels"]
        add_table(doc, labels + ["План", "Засчитано", "Брак", "Осталось", "%"],
                  [[r.get(k, "") for k in labels] + [r["План"], r["Засчитано"], r["Брак"], r["Осталось"],
                                                     f"{r['%']}%" if r.get("%") is not None else "—"]
                   for r in q["rows"] if not (filters or {}).get("city") or r.get("Город") == filters["city"]],
                  align_right=tuple(range(len(labels), len(labels) + 5)), font_size=8.5)
        n_sec += 1

    # ── Выводы ────────────────────────────────────────────────────────────
    heading(doc, f"{n_sec}. Выводы и рекомендации")
    if reasons:
        bullet(doc, f" «{engine.short_label(reasons[0][0])}» — {anket(reasons[0][1])}. "
                    f"{review.explain(cfg, reasons[0][0], '')[0]}", "Главная проблема:")
    reds = [r[1] for r in inter_rows if r[0] == STATUS_TEXT["RED"]]
    if reds:
        bullet(doc, f" {', '.join(map(str, reds))} — остановить, перепроверить анкеты (прозвон), решить о замене.",
               "Интервьюеры в красной зоне:")
    bad_cities = [c for c, _, n, _, p in city_stats if level(p, cfg) == "RED" and n >= 5]
    if bad_cities:
        bullet(doc, f" {', '.join(map(str, bad_cities))} — усилить контроль супервайзера.", "Города с высоким браком:")
    todo = int(((iv["is_defect"] | iv["is_warning"]) & ~iv["decision"].map(lambda v: isinstance(v, str))).sum())
    if todo:
        bullet(doc, f" {todo} {engine.plural(todo, 'подозрительная анкета', 'подозрительные анкеты', 'подозрительных анкет')} "
                    f"ещё без решения руководителя.", "Не разобрано:")
    if not (reasons or reds or bad_cities or todo):
        bullet(doc, " существенных нарушений нет.", "Итог:")
    para(doc, "Комментарий руководителя проекта", 11, NAVY, bold=True, space_after=2)
    box = doc.add_table(rows=1, cols=1)
    _borders(box, "A6A6A6")
    box.rows[0].cells[0].paragraphs[0].add_run("Впишите сюда выводы, договорённости с заказчиком и следующие шаги…").italic = True
    for _ in range(4):
        box.rows[0].cells[0].add_paragraph()
    doc.add_paragraph()

    # ── Приложение ────────────────────────────────────────────────────────
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    heading(doc, "Приложение. Анкеты с браком")
    brak = df[df["state"] == "brak"].sort_values(["city", "inter", "start"])
    rows = []
    for x in brak.head(600).itertuples():
        main = next(((c, t) for c, s, t in x.issues if c == x.primary), None)
        blk = next((b for (c, s, t), b in zip(x.issues, x.blocks) if c == x.primary), "")
        rows.append([x.row_id, x.city, x.inter, "" if pd.isna(x.start) else x.start.strftime("%d.%m %H:%M"), blk,
                     (engine.short_label(main[0]) + ": " + main[1]) if main else "",
                     x.decision if isinstance(x.decision, str) else ""])
    if rows:
        add_table(doc, ["ID", "Город", "Интервьюер", "Когда", "Блок", "Почему", "Решение"], rows,
                  widths=[1.6, 2.2, 2, 1.8, 2.4, 5.6, 1.6], font_size=7.5)
        if len(brak) > 600:
            para(doc, f"Показаны первые 600 из {len(brak)}. Полный список — в Excel-отчёте.", 9, GREY)
    else:
        para(doc, "Анкет с браком нет.")

    # номер страницы в нижнем колонтитуле
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = fp.add_run(f"{project} · контроль качества · стр. ")
    run.font.size, run.font.color.rgb = Pt(8), GREY
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        r = fp.add_run()
        r.font.size = Pt(8)
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r._r.append(el)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ─────────────────────────────────────────────────────────────────────────
# Карточки интервьюеров — для планёрки с супервайзером
# ─────────────────────────────────────────────────────────────────────────
# Что сказать интервьюеру по каждой причине — коротко и по делу.
ADVICE = {
    "too_short": "Не торопиться: зачитывать каждый вопрос полностью, дожидаться ответа.",
    "too_long": "Закрывать анкету сразу после интервью, не оставлять её открытой.",
    "start_gap": "Между интервью нужно время найти нового респондента — не открывать анкеты одну за другой.",
    "no_rest": "Начинать следующую анкету только с новым респондентом, а не сразу после предыдущей.",
    "overlap": "Открывать новую анкету только после того, как закончена предыдущая.",
    "mass_open": "Не открывать анкеты заранее «пачкой» — одна анкета на одного респондента.",
    "conveyor": "Слишком много анкет подряд: проверить, что каждое интервью реально проводится.",
    "night": "Работать только в рабочее время.",
    "dup_phone": "Не опрашивать одного человека дважды, записывать номер респондента, а не свой.",
    "dup_name": "Не опрашивать одного человека дважды.",
    "probe_depth": "Переспрашивать «А ещё?» столько раз, сколько требует инструкция.",
    "probe_low_avg": "Переспрашивать «А ещё?» в открытых вопросах — у вас ответов заметно меньше, чем у коллег.",
    "block_empty": "Не пропускать обязательные блоки анкеты.",
    "grid_same": "Зачитывать каждый пункт сетки, не отмечать один вариант во всех строках.",
    "near_dup": "Каждая анкета — отдельный респондент: анкеты не копировать и не заполнять по образцу.",
    "block_fast": "Не «пролетать» блоки: зачитывать вопросы и варианты ответа.",
    "no_gps": "Включать GPS на телефоне до начала интервью.",
    "geo_far": "Работать на своей точке опроса, в пределах радиуса.",
    "geo_city": "Проводить интервью в своём городе, указывать город правильно.",
    "geo_cluster": "Не делать слишком много анкет в одном месте — переходить по точкам маршрута.",
    "geo_same": "Координаты должны быть настоящими: не подменять GPS, не копировать анкеты.",
    "geo_jump": "Проверить время на телефоне и GPS: перемещения между анкетами невозможны по скорости.",
    "no_device": "Работать только с выданного телефона с включённым Device ID.",
    "device_multi_inter": "Каждый интервьюер работает со своего телефона и под своим кодом.",
    "inter_multi_device": "Работать под своим кодом с одного телефона.",
    "external": "Группа мониторинга забраковала интервью по аудиозаписи — разобрать эти записи вместе.",
}


def _new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(10.5)
    return doc


def cards(result, cfg, decisions_by_pos, project, inter=None):
    """Карточка на каждого интервьюера (или на одного): статус, цифры,
    причины брака с советом, работа по дням, необычные ответы, анкеты."""
    from . import daily
    df = frame(result, decisions_by_pos)
    iv_all = df[df["state"] != "tech"]
    inters = [inter] if inter else sorted(iv_all["inter"].dropna().unique(), key=str)
    d_rows, d_days, _, norm = daily.table(result, cfg, decisions_by_pos)
    d_by = {r["inter"]: r for r in d_rows}
    pats = {p["inter"]: p for p in result.get("answer_patterns") or []}
    doc = _new_doc()
    first = True
    for name in inters:
        g = iv_all[iv_all["inter"] == name]
        if g.empty:
            continue
        if not first:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        first = False
        n = len(g)
        b = int((g["state"] == "brak").sum())
        w = int((g["state"] == "warn").sum())
        pct = b / n * 100 if n else 0
        lvl = level(pct, cfg)
        para(doc, f"КАРТОЧКА ИНТЕРВЬЮЕРА · {project}", 9, GREY, bold=True, space_after=2)
        para(doc, str(name), 22, NAVY, bold=True, space_after=2)
        para(doc, f"{g['city'].mode().iloc[0] if len(g['city'].dropna()) else '—'} · {engine.data_period(g) or ''}", 10, GREY)
        para(doc, f"{STATUS_TEXT[lvl]} — брак {pct:.0f}%", 14, bold=True, space_after=4)
        para(doc, engine.STATUS_LEVELS[lvl]["actions"].capitalize() + ".", 10, GREY)
        dr = d_by.get(name)
        add_table(doc, ["Анкет", "Брак", "Проверить", "Норма", "Дней работал", "В среднем в день"],
                  [[n, b, w, n - b - w, dr["days_worked"] if dr else "—", dr["avg"] if dr else "—"]],
                  align_right=(0, 1, 2, 3, 4, 5), font_size=10)
        reasons = reason_counts(g)
        if reasons:
            para(doc, "Почему брак и что исправить", 12, NAVY, bold=True, space_after=2)
            add_table(doc, ["Причина", "Анкет", "Что исправить"],
                      [[engine.short_label(c), k, ADVICE.get(c, "Разобрать эти анкеты с супервайзером.")]
                       for c, k in reasons[:8]], widths=[4.5, 1.5, 11], align_right=(1,), font_size=9)
        warn = Counter(c for xs in g.loc[g["state"] == "warn", "issues"] for c in {c for c, s, _ in xs if s == engine.WARNING})
        if warn:
            para(doc, "Обратить внимание (не брак, но проверить)", 11, NAVY, bold=True, space_after=2)
            for c, k in warn.most_common(5):
                bullet(doc, f" — {k} анк. {ADVICE.get(c, '')}", engine.short_label(c))
        if dr and d_days:
            para(doc, "Работа по дням" + (f" (норма {norm} засчитанных)" if norm else ""), 11, NAVY, bold=True, space_after=2)
            add_table(doc, ["День", "Засчитано", "Брак", "Итог"],
                      [[pd.Timestamp(d).strftime("%d.%m"), c["ok"] if c["n"] else "—", c["brak"] or "—",
                        "не работал" if not c["n"] else ("ниже нормы" if c["status"] == "YELLOW" else "норма")]
                       for d, c in zip(d_days, dr["cells"])], align_right=(1, 2), font_size=9)
        if name in pats:
            para(doc, "Необычные ответы его респондентов", 11, NAVY, bold=True, space_after=2)
            for t in pats[name]["notes"]:
                bullet(doc, " " + t)
        worst = g[g["state"] == "brak"].sort_values("risk", ascending=False).head(10)
        if len(worst):
            para(doc, "Анкеты для разбора", 11, NAVY, bold=True, space_after=2)
            add_table(doc, ["ID", "Когда", "Почему"],
                      [[x.row_id, "" if pd.isna(x.start) else x.start.strftime("%d.%m %H:%M"),
                        "; ".join(t for c, s, t in x.issues if s == engine.DEFECT)[:220]] for x in worst.itertuples()],
                      widths=[2, 2.2, 12.8], font_size=8.5)
        para(doc, "Подпись интервьюера: ____________    Супервайзер: ____________    Дата: ________", 10, space_after=0)
    if first:
        para(doc, "Нет анкет для карточек.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
