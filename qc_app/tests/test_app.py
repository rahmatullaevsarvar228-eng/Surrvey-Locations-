# -*- coding: utf-8 -*-
import io
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))

from anketa_qc.remote import RemoteClient, RemoteError, combine, rows_to_frame  # noqa: E402
from anketa_qc.server import create_app  # noqa: E402
import make_demo  # noqa: E402


@pytest.fixture(scope="module")
def demo_bytes(tmp_path_factory):
    path = tmp_path_factory.mktemp("demo") / "demo.xlsx"
    make_demo.main(str(path))
    return path.read_bytes()


class FakeClient:
    """Сервер доступа в памяти — тот же протокол, что у server/Code.gs."""
    users = {"admin": ("pass-admin", "admin", True), "ali": ("pass-ali", "user", True)}
    frames = {}

    def __init__(self, url):
        if not url.startswith("https://"):
            raise RemoteError("Адрес сервера должен начинаться с https://")
        self.user = None
        self.calls = []

    def login(self, login, password):
        u = self.users.get(login)
        if not u or u[0] != password:
            raise RemoteError("Неверный логин или пароль")
        self.user = login
        return {"token": "t", "user": {"login": login, "name": login, "role": u[1]},
                "sources": [{"id": k, "name": k} for k in self.frames]}

    def _check(self):
        if not self.users.get(self.user, (None, None, False))[2]:
            raise RemoteError("доступ закрыт администратором", auth=True)

    def call(self, action, **params):
        self._check()
        self.calls.append((action, params))
        if action == "sources":
            return {"sources": [{"id": k, "name": k} for k in self.frames]}
        if action == "add_source":
            if "docs.google.com" not in params["url"]:
                raise RemoteError("Нужна ссылка на Google-таблицу")
            self.frames[params["name"]] = pd.DataFrame({"x": [1]})
            return {"source": {"id": params["name"]}}
        if action == "delete_source":
            self.frames.pop(params["id"], None)
            return {"ok": True}
        if action == "create_user":
            return {"login": params["login"], "password": "Abc123xyz9"}
        return {"ok": True}

    def fetch_frame(self, source_id):
        self._check()
        return source_id, self.frames[source_id]


def login(client, who="admin"):
    res = client.post("/api/auth/login", json={"server_url": "https://script.google.com/x/exec",
                                               "login": who, "password": f"pass-{who}"})
    assert res.status_code == 200, res.get_json()
    return res.get_json()


def test_rows_to_frame_and_combine():
    df = rows_to_frame(["a", "b", "a"], [[1, "x", 2], ["", "y"]])
    assert list(df.columns) == ["a", "b", "a.1"] and df["a"].isna().iloc[1]
    sheets = combine([("T1", df), ("T2", df.iloc[:1])])
    first = next(iter(sheets))
    assert first == "Все источники (2)" and len(sheets[first]) == 3
    assert list(sheets[first]["Источник"]) == ["T1", "T1", "T2"]
    assert set(sheets) == {first, "T1", "T2"}
    assert list(combine([("T1", df)])) == ["T1"]


def test_remote_client_maps_auth_errors(monkeypatch):
    import json as _json

    class Resp:
        def __init__(self, body):
            self.body = body

        def read(self):
            return _json.dumps(self.body).encode()

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    from anketa_qc import remote
    answers = iter([{"error": "AUTH: доступ закрыт администратором"}, {"error": "Нет доступа"},
                    b"<html>"])
    def fake_urlopen(req, timeout):
        a = next(answers)
        return Resp(a) if isinstance(a, dict) else type("R", (Resp,), {"read": lambda self: a})(None)
    monkeypatch.setattr(remote.urllib.request, "urlopen", fake_urlopen)
    c = RemoteClient("https://script.google.com/x/exec")
    with pytest.raises(RemoteError) as e:
        c.call("me")
    assert e.value.auth and "закрыт" in str(e.value)
    with pytest.raises(RemoteError) as e:
        c.call("me")
    assert not e.value.auth
    with pytest.raises(RemoteError, match="не то"):
        c.call("me")
    with pytest.raises(RemoteError, match="https"):
        RemoteClient("http://evil")


def test_login_required(tmp_path):
    client = create_app(tmp_path, client_factory=FakeClient).test_client()
    assert client.get("/").status_code == 200
    assert client.get("/api/auth").get_json()["logged_in"] is False
    res = client.get("/api/state")
    assert res.status_code == 401 and res.get_json()["auth"]
    bad = client.post("/api/auth/login", json={"server_url": "https://s/exec", "login": "ali", "password": "x"})
    assert bad.status_code == 400 and "Неверный" in bad.get_json()["error"]
    login(client, "ali")
    assert client.get("/api/state").status_code == 200
    # обычный сотрудник не может в администрирование
    assert client.post("/api/admin/list_users", json={}).status_code == 403
    # адрес сервера и логин запоминаются, пароль — нет
    auth = client.get("/api/auth").get_json()
    assert auth["server_url"].startswith("https://") and auth["last_login"] == "ali"
    assert "pass" not in str(auth)


def test_blocked_user_is_logged_out(tmp_path, demo_bytes):
    FakeClient.frames = {"Ташкент": pd.read_excel(io.BytesIO(demo_bytes), sheet_name="data")}
    client = create_app(tmp_path, client_factory=FakeClient).test_client()
    login(client, "ali")
    FakeClient.users = dict(FakeClient.users, ali=("pass-ali", "user", False))
    try:
        res = client.post("/api/source/remote", json={"ids": ["Ташкент"]})
        assert res.status_code == 401 and "закрыт" in res.get_json()["error"]
        assert client.get("/api/state").status_code == 401   # данные из памяти недоступны
    finally:
        FakeClient.users = dict(FakeClient.users, ali=("pass-ali", "user", True))


def test_admin_proxy(tmp_path):
    client = create_app(tmp_path, client_factory=FakeClient).test_client()
    login(client, "admin")
    res = client.post("/api/admin/create_user", json={"login": "new", "team": "Uzum"}).get_json()
    assert res["password"] == "Abc123xyz9"
    assert client.post("/api/admin/drop_database", json={}).status_code == 404


def test_team_adds_and_removes_own_sources(tmp_path):
    FakeClient.frames = {}
    client = create_app(tmp_path, client_factory=FakeClient).test_client()
    login(client, "ali")
    bad = client.post("/api/remote/sources/add", json={"name": "X", "url": "https://evil"})
    assert bad.status_code == 400
    lst = client.post("/api/remote/sources/add", json={"name": "Ташкент", "url": "https://docs.google.com/spreadsheets/d/A"}).get_json()
    assert [x["name"] for x in lst] == ["Ташкент"]
    assert client.post("/api/remote/sources/delete", json={"id": "Ташкент"}).get_json() == []


def test_auto_refresh_reports_new_anketas(tmp_path, demo_bytes):
    data = pd.read_excel(io.BytesIO(demo_bytes), sheet_name="data")
    FakeClient.frames = {"Поток": data.iloc[:200]}
    client = create_app(tmp_path, client_factory=FakeClient).test_client()
    login(client)
    assert client.post("/api/refresh").status_code == 400          # таблицы ещё не подключены
    state = client.post("/api/source/remote", json={"ids": ["Поток"]}).get_json()
    first = client.post("/api/refresh").get_json()   # колонки подставились сами
    assert first["summary"]["total"] == 200 and first["refresh"]["new"] == 0
    FakeClient.frames["Поток"] = data                                # в таблицу пришли новые анкеты
    second = client.post("/api/refresh").get_json()
    assert second["summary"]["total"] == len(data)
    assert second["refresh"]["new"] == len(data) - 200
    assert second["refresh"]["new_defects"] >= 1


def test_full_flow_through_api(tmp_path, demo_bytes):
    client = create_app(tmp_path, client_factory=FakeClient).test_client()
    login(client)
    data = pd.read_excel(io.BytesIO(demo_bytes), sheet_name="data")
    half = len(data) // 2
    FakeClient.frames = {"Волна — Ташкент": data.iloc[:half], "Волна — регионы": data.iloc[half:]}
    state = client.post("/api/source/remote", json={"ids": list(FakeClient.frames)}).get_json()
    assert state["config"]["sheet"] == "Все источники (2)"
    assert state["columns"][0] == "Источник"
    cfg = state["config"]
    cfg["mapping"] = state["suggested"]
    cols = state["columns"]
    first = next(i for i, c in enumerate(cols) if "первым приходит" in c)
    cfg["probing"]["blocks"] = [{"label": "Банки", "columns": cols[first:first + 7], "min_n": 4}]
    cfg["repetition"]["columns"] = cols[first:first + 7]
    cfg["repetition"]["values"] = [{"name": "Uzum Bank", "variants": ["uzum", "узум"]}]

    res = client.post("/api/run", json={"config": cfg})
    assert res.status_code == 200, res.get_json()
    r = res.get_json()
    assert r["summary"]["total"] > 300 and r["summary"]["defects"] > 0
    red = {x["Интервьюер"] for x in r["interviewers"] if x["Статус"] == "RED"}
    assert {"Inter 03", "Inter 08"} <= red
    assert r["repetition"]["enabled"]

    who = client.get("/api/answers/who", query_string={"answer": "Uzum Bank"}).get_json()
    assert who and {"ID анкеты", "Интервьюер"} <= set(who[0])

    data = client.get("/api/export/full").data
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["Интервьюеры"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert ws["A1"].fill.fgColor.rgb.endswith("1F4E78")

    hist = client.post("/api/history/save", json={"wave": "W1"}).get_json()
    assert hist["waves"][0]["wave"] == "W1"
    # настройки проекта переживают перезапуск приложения
    client2 = create_app(tmp_path, client_factory=FakeClient).test_client()
    login(client2)
    state2 = client2.get("/api/state").get_json()
    assert state2["config"]["remote_sources"] == list(FakeClient.frames)
    assert state2["config"]["repetition"]["values"][0]["name"] == "Uzum Bank"
    assert state2["waves"][0]["wave"] == "W1"


def test_run_requires_columns(tmp_path, demo_bytes):
    client = create_app(tmp_path, client_factory=FakeClient).test_client()
    login(client)
    client.post("/api/source/file", data={"file": (io.BytesIO(demo_bytes), "demo.xlsx")},
                content_type="multipart/form-data")
    # обязательные колонки, которых нет среди подсказок, — понятная ошибка
    state = client.get("/api/state").get_json()
    cfg = state["config"]
    cfg["mapping"]["device"] = None
    client.post("/api/config", json={"config": cfg})
    client.post("/api/sheet", json={"sheet": "Свод"})
    res = client.post("/api/run", json={})
    assert res.status_code == 400 and "Укажите колонки" in res.get_json()["error"]
    # а на листе с анкетами подсказки подставляются сами, без захода в «Колонки»
    client.post("/api/sheet", json={"sheet": "data"})
    assert client.post("/api/run", json={}).status_code == 200


def test_large_file_memory_friendly(tmp_path):
    """1500 анкет × 20 открытых колонок: общие данные (индекс ответов) не
    копируются на каждую строку, прогон укладывается в секунды."""
    import time
    from anketa_qc import config, engine
    n = 1500
    df = pd.DataFrame({
        "deviceid": [f"d{i % 60}" for i in range(n)], "inter": [f"I{i % 60}" for i in range(n)],
        "city": [f"C{i % 7}" for i in range(n)],
        "start": pd.date_range("2025-09-01 08:00", periods=n, freq="7min").astype(str),
        "end": (pd.date_range("2025-09-01 08:00", periods=n, freq="7min") + pd.Timedelta(minutes=6)).astype(str),
        **{f"q{j}": [f"бренд {(i * j) % 13}" for i in range(n)] for j in range(20)},
    })
    cfg = config.default_config()
    cfg["mapping"].update(device="deviceid", inter="inter", city="city", start="start", end="end")
    cfg["probing"]["blocks"] = [{"label": "Q", "columns": [f"q{j}" for j in range(20)], "min_n": 3}]
    t = time.time()
    res = engine.run(df, cfg)
    assert time.time() - t < 30
    assert "issues" in res["df"] and "answers" not in res["df"].columns


def test_access_server_script():
    """server/Code.gs в имитации сервисов Google (tests/gas_harness.js)."""
    import shutil
    import subprocess
    node = shutil.which("node")
    if not node:
        pytest.skip("нет Node.js")
    harness = Path(__file__).resolve().parent / "gas_harness.js"
    res = subprocess.run([node, str(harness)], capture_output=True, text=True, encoding="utf-8")
    assert res.returncode == 0, res.stdout + res.stderr


def test_gps_through_api(tmp_path, demo_bytes):
    client = create_app(tmp_path, client_factory=FakeClient).test_client()
    login(client)
    state = client.post("/api/source/file", data={"file": (io.BytesIO(demo_bytes), "demo.xlsx")},
                        content_type="multipart/form-data").get_json()
    assert state["suggested"]["lat"] == "_Координаты_latitude"
    cfg = state["config"]
    cfg["mapping"] = state["suggested"]
    cfg["geo"]["plan"] = client.get("/api/geo/default-plan").get_json()
    r = client.post("/api/run", json={"config": cfg}).get_json()
    g = r["geo"]
    assert g["enabled"] and g["far"] > 0 and g["clusters_over"] > 0 and g["same"] > 0
    flagged = {p["inter"] for p in g["points"] if p["far"]}
    assert "Inter 08" in flagged                     # «сидит дома»
    assert any(c["Интервьюер"] == "Inter 03" and c["Статус"] == "RED" for c in g["clusters"])
    wb = openpyxl.load_workbook(io.BytesIO(client.get("/api/export/full").data))
    assert {"GPS", "GPS скопления"} <= set(wb.sheetnames)
    # план из Excel
    buf = io.BytesIO()
    pd.DataFrame({"Город": ["Бухара"], "Широта": [39.77], "Долгота": [64.44]}).to_excel(buf, index=False)
    plan = client.post("/api/geo/plan/parse", data={"file": (io.BytesIO(buf.getvalue()), "plan.xlsx")},
                       content_type="multipart/form-data").get_json()
    assert plan["Бухара"]["points"][0]["lat"] == 39.77
