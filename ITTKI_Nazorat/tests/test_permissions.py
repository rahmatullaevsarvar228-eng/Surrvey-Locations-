"""Role permissions (every row of the table × every role), audit time in Tashkent, empty
'expected net benefit' field, E9 default threshold."""
import os
import sqlite3
from datetime import date as d, datetime, timezone

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from ittki import clock
from ittki.core.appsettings import AppSettings
from ittki.core.controls import ControlSettings
from ittki.db.repository import CostRecord, ImpairmentTest, ProjectRecord, Repository, User
from ittki.permissions import PERMISSIONS, PermissionDenied, ReadOnlyRole
from ittki.reports.excel_import import write_template
from ittki.services import REPORT_KINDS, DemoRefused, Service
from tests.test_db import all_yes, save_s2
from tests.test_engine import M

LEAD = User("Aliyev B.", "leader")
ACC = User("Karimova D.", "accountant")
AUD = User("Usmonov R.", "auditor")
USERS = {"leader": LEAD, "accountant": ACC, "auditor": AUD}

# The specification table: row -> roles allowed (✓)
TABLE = {
    "view":       {"leader", "accountant", "auditor"},   # all windows, tabs, audit log
    "project":    {"leader", "accountant"},              # create a project, save the passport
    "stage":      {"leader"},
    "assessment": {"leader"},                            # incl. draft
    "costs":      {"accountant"},                        # add, edit, delete, import
    "impairment": {"accountant"},
    "settings":   {"accountant"},
    "demo":       {"accountant"},                        # only into an empty database
    "reports":    {"leader", "accountant", "auditor"},   # all reports and exports
    "backup":     {"leader", "accountant", "auditor"},
}


@pytest.fixture(autouse=True)
def real_clock():
    yield
    clock.set_fixed(None)
    clock.set_now(None)


@pytest.fixture
def svc(tmp_path):
    repo = Repository(tmp_path / "p.sqlite3")
    pid = save_s2(repo)
    s = Service(repo, today=lambda: d(2026, 10, 11))
    s.pid = pid
    s.tmp = tmp_path
    s.demo = Service(Repository(tmp_path / "empty.sqlite3"))       # demo loads only into an empty database
    yield s
    s.demo.repo.close()
    repo.close()


def denied(repo):
    return [x for x in repo.audit_log() if x.action == "denied"]


def snapshot(repo):
    """Everything an action could change, except the audit log itself."""
    q = lambda sql: repo.conn.execute(sql).fetchall()
    return [tuple(map(tuple, q(f"SELECT * FROM {t} ORDER BY id"))) for t in
            ("projects", "stage_events", "assessments", "costs", "impairment_events", "impairment_tests")] + \
           [tuple(map(tuple, q("SELECT * FROM settings ORDER BY key")))]


# ------------------------------------------------------------------ the table itself
def test_permission_table_is_exactly_the_specification():
    rows = {"view": "view", "project": "project", "stage": "stage", "assessment": "assessment",
            "costs": "cost", "impairment": "impairment", "settings": "settings", "demo": "demo",
            "reports": "report", "backup": "backup"}
    for row, action in rows.items():
        assert set(PERMISSIONS[action]) == TABLE[row], row
    assert set(PERMISSIONS["import"]) == TABLE["costs"]          # import is part of the 'costs' row
    assert set(PERMISSIONS) == set(rows.values()) | {"import"}


# ------------------------------------------------------------------ real actions through the service layer
def _do_view(s, u):
    s.authorize(u, "view")
    assert s.analyze_all() and s.repo.audit_log() and s.findings_all() is not None


def _do_project(s, u):
    pid = s.save_project(u, ProjectRecord("NEW-1", "Yangi"))
    rec = s.repo.project(s.pid)
    rec.note = "pasport"
    s.save_project(u, rec)
    assert s.repo.project(pid).code == "NEW-1"


def _do_stage(s, u):
    s.change_stage(u, s.pid, "abandoned", d(2026, 1, 10), "test")
    assert s.repo.project(s.pid).abandoned_on == d(2026, 1, 10)


def _do_assessment(s, u):
    s.add_assessment(u, all_yes(s.pid, d(2026, 1, 5)))
    draft = all_yes(s.pid, d(2026, 1, 6))
    draft.answers["K5_adequate_resources"].evidence = None
    draft.is_draft = True
    s.add_assessment(u, draft)


def _do_costs(s, u):
    cid = s.add_cost(u, CostRecord(s.pid, d(2026, 1, 5), 7 * M, "salary"))
    s.update_cost(u, CostRecord(s.pid, d(2026, 1, 5), 8 * M, "salary", id=cid))
    s.delete_cost(u, cid)
    tpl = s.tmp / "import.xlsx"
    write_template(tpl)
    from openpyxl import load_workbook
    wb = load_workbook(tpl)
    wb["Xarajatlar"].append([d(2026, 2, 1), 3 * M, "materials", "PD", ""])
    wb.save(tpl)
    recs, errors = s.read_import(u, s.pid, tpl)
    assert s.import_costs(u, recs) == 1 and errors == []


def _do_impairment(s, u):
    s.set_impairment_event(u, s.pid, d(2025, 12, 31), 300 * M)
    tid = s.add_impairment_test(u, ImpairmentTest(s.pid, d(2025, 12, 20)))
    s.delete_impairment_test(u, tid)


def _do_settings(s, u):
    st = s.settings
    st.controls = ControlSettings(e9_amount=2_000_000)
    s.save_settings(u, st)


def _do_demo(s, u):
    assert s.demo.load_demo(u) == 4


def _do_reports(s, u):
    for kind in REPORT_KINDS:
        out = s.tmp / f"r_{kind}.{'pdf' if kind == 'audit_pdf' else 'docx' if kind == 'note' else 'xlsx'}"
        s.export(u, kind, out, as_of=d(2026, 10, 11), start=d(2025, 1, 1), end=d(2025, 12, 31))
        assert out.exists()


def _do_backup(s, u):
    s.backup(u, s.tmp / "zaxira.sqlite3")
    assert (s.tmp / "zaxira.sqlite3").exists()


ACTIONS = {"view": _do_view, "project": _do_project, "stage": _do_stage, "assessment": _do_assessment,
           "costs": _do_costs, "impairment": _do_impairment, "settings": _do_settings, "demo": _do_demo,
           "reports": _do_reports, "backup": _do_backup}
ROW_ACTION = {"costs": "cost", "reports": "report"}


@pytest.mark.parametrize("role", ["leader", "accountant", "auditor"])
@pytest.mark.parametrize("row", list(TABLE))
def test_every_row_for_every_role(svc, row, role):
    user = USERS[role]
    repo = svc.demo.repo if row == "demo" else svc.repo
    before = snapshot(repo)
    if role in TABLE[row]:
        ACTIONS[row](svc, user)
        assert denied(repo) == []
    else:
        exc = ReadOnlyRole if role == "auditor" else PermissionDenied
        with pytest.raises(exc) as e:
            ACTIONS[row](svc, user)
        assert str(e.value).startswith("Ruxsat yo'q")
        (rec,) = denied(repo)
        assert (rec.entity, rec.actor, rec.role) == (ROW_ACTION.get(row, row), user.name, role)
        assert snapshot(repo) == before                          # nothing changed


def test_demo_only_into_empty_database(svc):
    assert svc.repo.projects() and not svc.can_load_demo(ACC)
    with pytest.raises(DemoRefused):
        svc.load_demo(ACC)
    empty = svc.demo
    assert empty.can_load_demo(ACC) and not empty.can_load_demo(LEAD) and not empty.can_load_demo(AUD)
    empty.load_demo(ACC)
    assert not empty.can_load_demo(ACC)
    assert empty.repo.audit_log("demo")[0].actor == ACC.name


def test_repository_is_guarded_by_the_same_check(svc):
    """Direct repository calls (bypassing the service) are refused the same way."""
    with pytest.raises(PermissionDenied):
        svc.repo.add_cost(CostRecord(svc.pid, d(2024, 1, 1), 1, "salary"), LEAD)
    with pytest.raises(PermissionDenied):
        svc.repo.change_stage(svc.pid, "abandoned", d(2026, 1, 1), ACC)
    assert [x.entity for x in denied(svc.repo)] == ["cost", "stage"]


def test_denied_record_survives_reopen_and_cannot_be_deleted(svc):
    with pytest.raises(PermissionDenied):
        svc.add_cost(LEAD, CostRecord(svc.pid, d(2024, 1, 1), 1, "salary"))
    path = svc.repo.path
    svc.repo.close()
    r2 = Repository(path)
    (row,) = [x for x in r2.audit_log() if x.action == "denied"]
    with pytest.raises(sqlite3.DatabaseError):
        r2.conn.execute("DELETE FROM audit_log WHERE id = ?", (row.id,))
    r2.close()
    svc.repo = Repository(path)


# ------------------------------------------------------------------ UI: buttons and refusal messages
@pytest.fixture
def app():
    pytest.importorskip("PySide6.QtWidgets")
    from ittki.ui.app import make_app
    return make_app(["t"])


def _buttons(svc, user):
    """Row -> list of (button, expected_enabled_if_allowed)."""
    from ittki.ui.main_window import MainWindow
    from ittki.ui.project_card import ProjectCard
    w = MainWindow(svc, user)
    for i in range(len(w.pages)):
        w.show_page(i)
    home, projects, controls, reports, settings, audit = w.pages
    card = ProjectCard(svc, user, svc.pid)
    rep_buttons = [b for b in reports.findChildren(type(card.b_save)) if b.text() and "(" in b.text()]
    return w, card, {
        "project": [projects.b_new, card.b_save],
        "stage": [card.b_stage],
        "assessment": [card.b_assess],
        "costs": [card.b_c_add, card.b_c_edit, card.b_c_del, card.b_c_imp],
        "impairment": [card.b_imp, card.b_test_add, card.b_test_del],
        "settings": [settings.b_save],
        "reports": rep_buttons + [card.b_am_exp, audit.b_export],
        "backup": [home.b_backup],
    }, home


@pytest.mark.parametrize("role", ["leader", "accountant", "auditor"])
def test_buttons_follow_the_table(app, svc, role):
    user = USERS[role]
    w, card, rows, home = _buttons(svc, user)
    assert len(rows["reports"]) == 10
    for row, buttons in rows.items():
        for b in buttons:
            assert b.isEnabled() == (role in TABLE[row]), (row, b.text())
    assert not home.b_demo.isEnabled()                    # database is not empty -> nobody
    for i in range(card.tabs.count()):                    # view: every tab opens for every role
        card.tabs.setCurrentIndex(i)
    card.close(); w.close()


@pytest.mark.parametrize("role", ["leader", "accountant", "auditor"])
def test_demo_button_on_empty_database(app, tmp_path, role):
    from ittki.ui.pages import HomePage
    s = Service(Repository(tmp_path / "empty.sqlite3"))
    page = HomePage(s, USERS[role])
    page.refresh()
    assert page.b_demo.isEnabled() == (role == "accountant")
    s.repo.close()


@pytest.mark.parametrize("role,handler,action", [
    ("leader", "_add_cost", "cost"), ("leader", "_import_costs", "import"), ("leader", "_add_test", "impairment"),
    ("accountant", "_change_stage", "stage"), ("accountant", "_new_assessment", "assessment"),
    ("auditor", "_save_info", "project"), ("auditor", "_del_cost", "cost")])
def test_forbidden_action_called_anyway_shows_message_and_is_audited(app, svc, monkeypatch, role, handler, action):
    from ittki.ui import project_card
    msgs = []
    monkeypatch.setattr("ittki.ui.widgets.error", lambda parent, text: msgs.append(text))
    for dlg in ("CostDialog", "AssessmentDialog", "StageDialog"):
        monkeypatch.setattr(project_card, dlg, lambda *a, **k: pytest.fail("dialog must not open"))
    card = project_card.ProjectCard(svc, USERS[role], svc.pid)
    before = snapshot(svc.repo)
    getattr(card, handler)()
    assert len(msgs) == 1 and msgs[0].startswith("Ruxsat yo'q")
    (rec,) = denied(svc.repo)
    assert rec.entity == action and rec.role == role
    assert snapshot(svc.repo) == before
    card.close()


def test_settings_page_and_demo_refused_for_leader(app, svc, monkeypatch):
    from ittki.ui.pages import HomePage, SettingsPage
    msgs = []
    monkeypatch.setattr("ittki.ui.widgets.error", lambda parent, text: msgs.append(text))
    page = SettingsPage(svc, LEAD)
    page.refresh()
    page.acc["expense"].setText("9430")
    page._save()
    HomePage(svc, LEAD)._demo()
    assert svc.settings.accounts.expense == "9420"
    assert [x.entity for x in denied(svc.repo)] == ["settings", "demo"] and len(msgs) == 2


def test_auditor_can_export_and_back_up_from_the_ui(app, svc, monkeypatch):
    from ittki.ui import pages
    out = svc.tmp / "rec.xlsx"
    monkeypatch.setattr(pages.QFileDialog, "getSaveFileName", lambda *a, **k: (str(out), ""))
    monkeypatch.setattr(pages, "info", lambda *a: None)
    page = pages.ReportsPage(svc, AUD)
    page._run("reconcile", "xlsx", "rep.reconcile")
    assert out.exists()
    out = svc.tmp / "z.sqlite3"
    pages.HomePage(svc, AUD)._backup()
    assert out.exists() and denied(svc.repo) == []


# ------------------------------------------------------------------ audit time: UTC in DB, Tashkent on screen
def test_audit_time_stored_utc_shown_in_tashkent(app, svc):
    from ittki.ui.pages import AuditPage
    clock.set_now(datetime(2026, 10, 12, 5, 41, 8, tzinfo=timezone.utc))
    svc.add_cost(ACC, CostRecord(svc.pid, d(2026, 1, 5), 1, "salary"))
    last = svc.repo.audit_log()[-1]
    assert last.at == "2026-10-12T05:41:08+00:00"                      # UTC in the database
    assert clock.fmt_tashkent(last.at) == "12.10.2026 10:41:08"
    page = AuditPage(svc, AUD)
    page.refresh()
    assert page.table.horizontalHeaderItem(0).text() == "Vaqt (Toshkent, UTC+5)"
    assert page.table.item(0, 0).text() == "12.10.2026 10:41:08"       # newest first
    out = svc.tmp / "audit.xlsx"
    svc.export(AUD, "audit_log", out)
    from openpyxl import load_workbook
    rows = list(load_workbook(out).active.iter_rows(values_only=True))
    head = next(i for i, r in enumerate(rows) if r[0] == "Vaqt (Toshkent, UTC+5)")
    assert rows[-1][0] == "12.10.2026 10:41:08" and head < len(rows) - 1


def test_time_conversion_rules():
    assert clock.fmt_tashkent("2026-10-12T05:41:08Z") == "12.10.2026 10:41:08"
    assert clock.fmt_tashkent("2026-10-12T23:30:00+00:00") == "13.10.2026 04:30:00"   # crosses midnight
    clock.set_fixed(d(2026, 10, 11))
    assert clock.today() == d(2026, 10, 11) and clock.utc_iso() == "2026-10-11T05:00:00+00:00"
    assert clock.utc_iso() == "2026-10-11T05:00:01+00:00"


def test_stage_history_shows_tashkent_time(app, svc):
    from ittki.ui.project_card import ProjectCard
    clock.set_now(datetime(2026, 10, 12, 5, 41, 8, tzinfo=timezone.utc))
    svc.change_stage(LEAD, svc.pid, "abandoned", d(2026, 1, 10))
    card = ProjectCard(svc, AUD, svc.pid)
    t = card.t_stages
    assert t.horizontalHeaderItem(6).text() == "Vaqt (Toshkent, UTC+5)"
    assert t.item(t.rowCount() - 1, 6).text() == "12.10.2026 10:41:08"   # one time per change
    card.close()


# ------------------------------------------------------------------ expected net benefit: empty vs 0
def test_money_edit_empty_is_none_not_zero(app):
    from ittki.ui.widgets import MoneyEdit
    e = MoneyEdit(None)
    assert e.text() == "" and e.placeholderText() == "" and e.value() is None
    e.setText("0")
    assert e.value() == 0


def test_card_keeps_benefit_empty_unless_zero_entered(app, svc):
    from ittki.ui.project_card import ProjectCard
    card = ProjectCard(svc, LEAD, svc.pid)
    assert card.f_benefit.text() == ""
    card._save_info()
    assert svc.repo.project(svc.pid).expected_net_benefit is None
    card.f_benefit.setText("0")
    card._save_info()
    assert svc.repo.project(svc.pid).expected_net_benefit == 0
    card.refresh()
    assert card.f_benefit.text() == "0"
    card.close()


# ------------------------------------------------------------------ E9 default
def test_e9_default_threshold_is_one_million(svc):
    assert ControlSettings().e9_amount == 1_000_000 == svc.settings.controls.e9_amount


def test_old_database_with_default_zero_is_upgraded(tmp_path):
    r = Repository(tmp_path / "o.sqlite3")
    r.conn.execute("UPDATE settings SET value = '0' WHERE key = 'control.e9_amount'")
    r.conn.execute("DELETE FROM settings WHERE key = 'meta.settings_version'")
    r.close()
    r = Repository(tmp_path / "o.sqlite3")
    assert r.settings().controls.e9_amount == 1_000_000
    assert r.audit_log("settings")[-1].after == {"control.e9_amount": 1_000_000}
    r.close()


def test_user_chosen_zero_is_kept(tmp_path):
    r = Repository(tmp_path / "u.sqlite3")
    s = r.settings()
    s.controls = ControlSettings(e9_amount=0)
    r.save_settings(s, ACC)
    r.conn.execute("DELETE FROM settings WHERE key = 'meta.settings_version'")
    r.close()
    r = Repository(tmp_path / "u.sqlite3")
    assert r.settings().controls.e9_amount == 0
    r.close()
