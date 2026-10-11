"""Stage 4: Excel / Word / PDF reports and Excel cost import."""
import random
from datetime import date as d

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from ittki.core.appsettings import AppSettings
from ittki.core.engine import carrying_amount
from ittki.db.repository import ProjectRecord, Repository, User
from ittki.demo import load_demo
from ittki.reports.excel import export_costs, export_findings, export_reconcile, export_schedule
from ittki.reports.excel_import import read_costs, write_template
from ittki.reports.pdf import export_audit_report
from ittki.reports.word import export_note
from ittki.services import Analysis, Service, movement
from tests.test_db import save_s2
from tests.test_engine import M
from tests.test_invariants import random_project

TODAY = d(2026, 10, 11)
USER = User("Raxmatullayev S.N.", "auditor")


@pytest.fixture
def svc(tmp_path):
    repo = Repository(tmp_path / "r.sqlite3")
    save_s2(repo)
    load_demo(repo)
    yield Service(repo, today=lambda: TODAY)
    repo.close()


def _cells(path, sheet=None):
    wb = load_workbook(path)
    ws = wb[sheet] if sheet else wb.active
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


def test_schedule_excel(svc, tmp_path):
    p = tmp_path / "s.xlsx"
    export_schedule(svc.analyze_all(), p, svc.settings)
    rows = _cells(p, "S2")
    total = next(r for r in rows if r[0] == "Jami")
    assert total[2] == 550 * M and total[6] == 550 * M                  # IAS 38 and BHMS 11 sums
    data = [r for r in rows if hasattr(r[0], "year")]
    assert len(data) == 96 and data[0][0].date() == d(2025, 1, 31)        # union of periods = 96 months
    assert data[59][8] == 0 and data[60][5] is None                       # BHMS 11 closes after 60 months


def test_reconcile_excel(svc, tmp_path):
    p = tmp_path / "r.xlsx"
    an = svc.analyze_all(d(2025, 12, 31))
    export_reconcile(an, p, svc.settings)
    rows = _cells(p)
    s2 = next(r for r in rows if r[0] == "S2")
    assert s2[6] == 550 * M - 12 * (550 * M // 96) and s2[7] == 550 * M - 12 * (550 * M // 60)
    assert s2[8] == s2[6] - s2[7] and "96 oy" in s2[9]
    atj = next(r for r in rows if r[0] == "DEMO-ATJBT-02")
    assert "K3" in atj[9]


def test_costs_excel(svc, tmp_path):
    p = tmp_path / "c.xlsx"
    export_costs(svc.analyze_all(), p, svc.settings)
    rows = [r for r in _cells(p) if r[0] == "S2"]
    assert len(rows) == 6 and sum(r[2] for r in rows) == 940 * M
    cap = [r for r in rows if r[6] == "Kapitallashtirish"]
    assert sum(r[7] for r in cap) == 550 * M and {r[8] for r in cap} == {"R7"}
    assert {r[10] for r in cap} == {"IAS 38.57"} and {r[15] for r in cap} == {"BHMS 11, 12-band"}


def test_costs_excel_r8_split(tmp_path):
    repo = Repository(tmp_path / "x.sqlite3")
    pid = save_s2(repo)
    p = repo.project(pid); p.expected_net_benefit = 400 * M; repo.save_project(p, User("L", "leader"))
    svc = Service(repo, today=lambda: TODAY)
    out = tmp_path / "c.xlsx"
    export_costs([svc.analyze(pid)], out, svc.settings)
    nov = [r for r in _cells(out) if r[0] == "S2" and hasattr(r[1], "month") and r[1].month == 11]
    assert [(r[11], r[12], r[13]) for r in nov] == [("Kapitallashtirish", 100 * M, "R7"), ("Xarajat", 150 * M, "R8")]
    assert nov[0][7] == 250 * M and nov[1][2] is None                     # IAS 38: whole cost capitalized
    repo.close()


def test_findings_excel_and_pdf(svc, tmp_path):
    an = svc.analyze_all()
    export_findings(an, tmp_path / "f.xlsx", svc.settings, USER, TODAY)
    rows = _cells(tmp_path / "f.xlsx")
    codes = {r[1] for r in rows if r[1] and str(r[1]).startswith("E")}
    assert {"E1", "E4", "E7", "E9", "E10"} <= codes
    pdf = tmp_path / "a.pdf"
    export_audit_report(an, pdf, svc.settings, USER, TODAY)
    data = pdf.read_bytes()
    assert data.startswith(b"%PDF") and len(data) > 3000


def test_word_note_numbers(svc, tmp_path):
    p = tmp_path / "n.docx"
    an = svc.analyze_all(d(2025, 12, 31))
    export_note([a for a in an if a.record.code == "S2"], p, svc.settings, d(2025, 1, 1), d(2025, 12, 31))
    text = "\n".join(par.text for par in Document(p).paragraphs)
    cells = "\n".join(c.text for t in Document(p).tables for row in t.rows for c in row.cells)
    assert "20-band" in text and "38.118" in text and "38.126" in text
    nas_close = 550 * M - 12 * (550 * M // 60)
    ias_close = 550 * M - 12 * (550 * M // 96)
    m = lambda v: f"{v:,}".replace(",", " ")
    assert m(nas_close) in cells and m(ias_close) in cells and m(550 * M) in cells
    # 2024: R&D expensed (IAS 38.126) = 100 (research) + 200 (criteria not met); training/admin are R1
    p2 = tmp_path / "n2.docx"
    export_note([a for a in an if a.record.code == "S2"], p2, svc.settings, d(2024, 1, 1), d(2024, 12, 31))
    text2 = "\n".join(par.text for par in Document(p2).paragraphs)
    assert f"{m(300 * M)} so'm" in text2 and m(90 * M) in text2
    assert "foydalanishga tayyor ITTKI natijalari yo'q" in text2          # asset ready only on 01.01.2025


def test_word_note_all_demo(svc, tmp_path):
    p = tmp_path / "n.docx"
    export_note(svc.analyze_all(), p, svc.settings, d(2026, 1, 1), d(2026, 9, 30))
    text = "\n".join(par.text for par in Document(p).paragraphs)
    assert "Shartli misol" in text


def test_movement_rollforward_identity_random():
    """opening + capitalized - amortization - impairment - write-offs - disposals = closing = engine carrying."""
    random.seed(11)
    acc = AppSettings().accounts
    from ittki.core import engine, journal as jr
    from ittki.db.repository import ProjectRecord as PR
    n = 0
    for _ in range(800):
        p = random_project()
        if p.scope != "rnd":
            continue
        for fw in ("IAS38", "NAS11"):
            res = engine.evaluate(p, fw, as_of=d(2026, 6, 30))
            a = Analysis(PR("x", "x"), p, d(2026, 6, 30), {fw: res}, {fw: engine.schedule(res, p)},
                         {fw: jr.journal(res, p)}, {}, {}, [], {})
            for start, end in ((d(2023, 1, 1), d(2023, 12, 31)), (d(2024, 1, 1), d(2024, 12, 31)),
                               (d(2025, 1, 1), d(2025, 12, 31))):
                m = movement(a, fw, start, end, acc)
                assert (m["carrying_open"] + m["capitalized"] - m["amortization"] - m["impairment"]
                        - m["wip_writeoff"] - m["disposal_residual"]) == m["carrying_close"]
                assert m["wip_open"] + m["capitalized"] - m["wip_writeoff"] - m["transfer"] == m["wip_close"]
                quirk = p.available_for_use and any(w >= p.available_for_use for w, _ in res["nas_writeoff_rows"])
                if not quirk:
                    assert m["carrying_close"] == carrying_amount(res, p, end)
                n += 1
    assert n > 3000


def test_template_roundtrip(tmp_path):
    tpl = tmp_path / "t.xlsx"
    write_template(tpl)
    wb = load_workbook(tpl)
    assert wb.sheetnames == ["Xarajatlar", "Toifalar", "Yo'riqnoma"]
    ws = wb["Xarajatlar"]
    ws.append([d(2025, 3, 14), 420_000_000, "salary", "PD-1", "ok"])
    ws.append(["15.04.2025", "1 000 000", "Materiallar", "PD-2", ""])
    ws.append([d(2025, 5, 1), 12.5, "salary", "", ""])            # fractional
    ws.append(["32.13.2025", 5, "salary", "", ""])                 # bad date
    ws.append([d(2025, 5, 1), 5, "bonus", "", ""])                 # unknown category
    ws.append([None, None, None, None, None])                      # empty -> skipped
    wb.save(tpl)
    recs, errors = read_costs(tpl, 7)
    assert [(r.on_date, r.amount, r.category, r.project_id) for r in recs] == [
        (d(2025, 3, 14), 420_000_000, "salary", 7), (d(2025, 4, 15), 1_000_000, "materials", 7)]
    assert len(errors) == 3 and all(e.startswith("Qator ") for e in errors)
    assert all(isinstance(r.amount, int) for r in recs)


def test_template_file_in_repo_is_current(tmp_path):
    from pathlib import Path
    shipped = Path(__file__).parent.parent / "templates" / "xarajatlar_shablon.xlsx"
    assert shipped.exists()
    fresh = tmp_path / "t.xlsx"
    write_template(fresh)
    a, b = load_workbook(shipped), load_workbook(fresh)
    assert a.sheetnames == b.sheetnames
    assert [c.value for c in a["Xarajatlar"][1]] == [c.value for c in b["Xarajatlar"][1]]
