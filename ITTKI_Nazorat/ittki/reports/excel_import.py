"""M4. Cost import from the Excel template (templates/xarajatlar_shablon.xlsx)."""
from __future__ import annotations

from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..core.engine import ALL_CATEGORIES, INCLUDED
from ..db.repository import CostRecord
from ..i18n import uz

COLUMNS = ["Sana (kk.oo.yyyy)", "Summa (so'm, butun son)", "Toifa", "Birlamchi hujjat", "Izoh"]
SHEET = "Xarajatlar"
_ORDER = sorted(INCLUDED["IAS38"]) + sorted(ALL_CATEGORIES - INCLUDED["IAS38"])


def write_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    for j, h in enumerate(COLUMNS, 1):
        c = ws.cell(1, j, h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EDF3")
    for r in range(2, 1002):
        ws.cell(r, 1).number_format = "DD.MM.YYYY"
        ws.cell(r, 2).number_format = "#,##0"
    for col, w in zip("ABCDE", (18, 24, 22, 20, 44)):
        ws.column_dimensions[col].width = w
    ref = wb.create_sheet("Toifalar")
    ref.append(["Kod", "Nomi", "IAS 38 aktivga", "BHMS 11 aktivga"])
    for c in ref[1]:
        c.font = Font(bold=True)
    for k in _ORDER:
        ref.append([k, uz.CATEGORY[k], "ha" if k in INCLUDED["IAS38"] else "yo'q",
                    "ha" if k in INCLUDED["NAS11"] else "yo'q"])
    ref.column_dimensions["A"].width = 22
    ref.column_dimensions["B"].width = 48
    dv = DataValidation(type="list", formula1=f"=Toifalar!$A$2:$A${len(_ORDER) + 1}", allow_blank=False)
    dv.error = "Toifa 'Toifalar' varag'idagi koddan bo'lishi kerak"
    ws.add_data_validation(dv)
    dv.add("C2:C5000")
    info = wb.create_sheet("Yo'riqnoma")
    for line in ("Xarajatlarni import qilish shabloni — ITTKI Nazorat",
                 "1. Har bir qator — bitta xarajat. Birinchi qator (sarlavha) o'zgartirilmaydi.",
                 "2. Sana: Excel sanasi yoki kk.oo.yyyy matni.",
                 "3. Summa: butun son, so'mda (kasr qism qabul qilinmaydi).",
                 "4. Toifa: 'Toifalar' varag'idagi kod (masalan, salary) yoki uning nomi.",
                 "5. Import: Loyiha kartasi → Xarajatlar → Exceldan import.",
                 "",
                 "Namuna: 14.03.2025 | 420 000 000 | salary | PD-001 | Tajriba qurilmasi loyihasi"):
        info.append([line])
    info.column_dimensions["A"].width = 90
    wb.save(path)


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                pass
    raise ValueError(f"sana noto'g'ri: {v!r}")


def _amount(v):
    if isinstance(v, bool):
        raise ValueError(f"summa noto'g'ri: {v!r}")
    if isinstance(v, int):
        out = v
    elif isinstance(v, float):
        if not v.is_integer():
            raise ValueError(f"summa butun son bo'lishi kerak: {v!r}")
        out = int(v)
    elif isinstance(v, str) and v.replace(" ", "").replace(" ", "").isdigit():
        out = int(v.replace(" ", "").replace(" ", ""))
    else:
        raise ValueError(f"summa noto'g'ri: {v!r}")
    if out <= 0:
        raise ValueError(f"summa musbat bo'lishi kerak: {v!r}")
    return out


_BY_LABEL = {v.lower(): k for k, v in uz.CATEGORY.items()}


def _category(v):
    s = str(v or "").strip()
    if s in ALL_CATEGORIES:
        return s
    if s.lower() in _BY_LABEL:
        return _BY_LABEL[s.lower()]
    raise ValueError(f"toifa noma'lum: {v!r}")


def read_costs(path, project_id: int) -> tuple[list[CostRecord], list[str]]:
    """Returns (valid records, error messages 'Qator N: ...'). Empty rows are skipped."""
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
    out, errors = [], []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        row = list(row) + [None] * 5
        if all(v in (None, "") for v in row[:5]):
            continue
        try:
            out.append(CostRecord(project_id, _date(row[0]), _amount(row[1]), _category(row[2]),
                                  str(row[3] or "").strip(), str(row[4] or "").strip()))
        except ValueError as e:
            errors.append(f"Qator {i}: {e}")
    return out, errors
