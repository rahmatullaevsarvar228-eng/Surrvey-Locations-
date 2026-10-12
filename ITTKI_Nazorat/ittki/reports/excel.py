"""Excel reports (openpyxl): amortization schedule (M7), NAS <-> IFRS reconciliation (M8),
cost journal with decisions (M4/M5), exception report (M10). Numbers come from the engine."""
from __future__ import annotations

from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .. import clock
from ..core.appsettings import AppSettings
from ..core.engine import FRAMEWORKS
from ..core.rules import reason_text, rule_code, rule_ref
from ..i18n import uz
from ..services import Analysis, cost_rows

MONEY = "#,##0"
DATE = "DD.MM.YYYY"
HEAD_FILL = PatternFill("solid", fgColor="E8EDF3")
TOTAL_FILL = PatternFill("solid", fgColor="F3F4F6")
THIN = Side(style="thin", color="C9CED6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREEN = Font(color="1E7B34", bold=True)
RED = Font(color="B3261E", bold=True)


def _title(ws, title: str, settings: AppSettings, sub: str = "", demo: bool = False) -> int:
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True, color="1F4E79")
    ws["A2"] = settings.organization
    ws["A3"] = sub
    row = 4
    if demo:
        ws["A4"] = uz.DEMO_MARK
        ws["A4"].font = Font(italic=True, color="8A6100")
        row = 5
    ws[f"A{row}"] = f"Tuzilgan: {clock.now_tashkent()} (Toshkent, UTC+5) · {uz.APP_NAME}"
    ws[f"A{row}"].font = Font(size=9, color="5F6368")
    return row + 2


def _header(ws, row: int, headers: list[str]):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row, j, h)
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = ws.cell(row + 1, 1)


def _put(ws, row: int, values: list, money_cols=(), date_cols=()):
    for j, v in enumerate(values, 1):
        c = ws.cell(row, j, v)
        c.border = BORDER
        if j in money_cols:
            c.number_format = MONEY
        if j in date_cols:
            c.number_format = DATE
    return row + 1


def _widths(ws, widths: list[int]):
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def _demo(analyses):
    return any(a.record.is_demo for a in analyses)


# ---------------------------------------------------------------- M7
def export_schedule(analyses: list[Analysis], path, settings: AppSettings):
    wb = Workbook()
    ws = wb.active
    ws.title = "Umumiy"
    r = _title(ws, "Amortizatsiya jadvali (IAS 38 | BHMS 11)", settings, "To'g'ri chiziqli usul, oy oxirida",
               _demo(analyses))
    heads = ["Loyiha", "Nomi", "Standart", "Amortizatsiya boshlanishi", "Muddat (oy)", "Amortizatsiya qilinadigan",
             "Jami amortizatsiya", "Jami qadrsizlanish"]
    _header(ws, r, heads)
    r += 1
    for a in analyses:
        for fw in FRAMEWORKS:
            sch = a.schedule[fw]
            if not sch:
                continue
            r = _put(ws, r, [a.record.code, a.record.name, uz.FW[fw], sch[0]["period_end"], a.res[fw]["amort_months"],
                             sch[0]["opening"], sum(x["amortization"] for x in sch),
                             sum(x["impairment"] for x in sch)], money_cols=(6, 7, 8), date_cols=(4,))
    _widths(ws, [16, 50, 10, 14, 10, 18, 18, 16])
    for a in analyses:
        if not any(a.schedule[fw] for fw in FRAMEWORKS):
            continue
        sh = wb.create_sheet(a.record.code[:31])
        r = _title(sh, f"{a.record.code} — {a.record.name}", settings, "Amortizatsiya jadvali", a.record.is_demo)
        sh.cell(r - 1, 2, "IAS 38").font = Font(bold=True)
        sh.cell(r - 1, 6, "BHMS 11").font = Font(bold=True)
        _header(sh, r, ["Davr", "Davr boshiga", "Amortizatsiya", "Qadrsizlanish", "Davr oxiriga",
                        "Davr boshiga", "Amortizatsiya", "Qadrsizlanish", "Davr oxiriga"])
        by = {fw: {x["period_end"]: x for x in a.schedule[fw]} for fw in FRAMEWORKS}
        periods = sorted(set(by["IAS38"]) | set(by["NAS11"]))
        r += 1
        for pe in periods:
            vals = [pe]
            for fw in FRAMEWORKS:
                x = by[fw].get(pe)
                vals += [x["opening"], x["amortization"], x["impairment"], x["closing"]] if x else [None] * 4
            r = _put(sh, r, vals, money_cols=range(2, 10), date_cols=(1,))
        tot = ["Jami"]
        for fw in FRAMEWORKS:
            tot += [None, sum(x["amortization"] for x in a.schedule[fw]), sum(x["impairment"] for x in a.schedule[fw]),
                    None]
        _put(sh, r, tot, money_cols=range(2, 10))
        for j in range(1, 10):
            sh.cell(r, j).font = Font(bold=True)
            sh.cell(r, j).fill = TOTAL_FILL
        _widths(sh, [12] + [16] * 8)
    wb.save(path)


# ---------------------------------------------------------------- M8
def export_reconcile(analyses: list[Analysis], path, settings: AppSettings):
    at = analyses[0].as_of if analyses else clock.today()
    wb = Workbook()
    ws = wb.active
    ws.title = "Solishtirma"
    r = _title(ws, "BHMS ↔ MHXS solishtirma (BHMS 11 (2024) | IAS 38)", settings,
               f"{uz.fdate(at)} holatiga", _demo(analyses))
    _header(ws, r, ["Loyiha", "Nomi", "Qaror IAS 38", "Qaror BHMS 11", "Kapitallashtirilgan IAS 38",
                    "Kapitallashtirilgan BHMS 11", "Balans qiymati IAS 38", "Balans qiymati BHMS 11", "Farq",
                    "Farq sabablari"])
    r += 1
    first = r
    for a in analyses:
        ri, rn = a.res["IAS38"], a.res["NAS11"]
        r = _put(ws, r, [a.record.code, a.record.name, uz.DECISION_SHORT[ri["decision"]],
                         uz.DECISION_SHORT[rn["decision"]], ri["capitalized"], rn["capitalized"],
                         a.carrying["IAS38"], a.carrying["NAS11"], a.carrying["IAS38"] - a.carrying["NAS11"],
                         "\n".join(a.reasons)], money_cols=(5, 6, 7, 8, 9))
        ws.cell(r - 1, 10).alignment = Alignment(wrap_text=True, vertical="top")
    tot = ["Jami", ""] + ["", ""] + [f"=SUM({get_column_letter(j)}{first}:{get_column_letter(j)}{r - 1})"
                                     for j in range(5, 10)] + [""]
    _put(ws, r, tot, money_cols=(5, 6, 7, 8, 9))
    for j in range(1, 11):
        ws.cell(r, j).font = Font(bold=True)
        ws.cell(r, j).fill = TOTAL_FILL
    _widths(ws, [16, 44, 18, 18, 18, 18, 18, 18, 16, 90])
    wb.save(path)


# ---------------------------------------------------------------- M4/M5
def export_costs(analyses: list[Analysis], path, settings: AppSettings):
    wb = Workbook()
    ws = wb.active
    ws.title = "Xarajatlar"
    r = _title(ws, "Xarajatlar jurnali va qarorlar (IAS 38 | BHMS 11)", settings,
               "Har bir xarajat: qaror, sabab, standart bandi", _demo(analyses))
    _header(ws, r, ["Loyiha", "Sana", "Summa", "Toifa", "Birlamchi hujjat", "Izoh",
                    "IAS 38: qaror", "IAS 38: summa", "IAS 38: qoida", "IAS 38: sabab", "IAS 38: band",
                    "BHMS 11: qaror", "BHMS 11: summa", "BHMS 11: qoida", "BHMS 11: sabab", "BHMS 11: band"])
    head = r
    r += 1
    for a in analyses:
        per = {fw: cost_rows(a, fw) for fw in FRAMEWORKS}
        for i, c in enumerate(a.costs):
            parts = {fw: per[fw][i] if i < len(per[fw]) else [] for fw in FRAMEWORKS}
            n = max(1, *(len(v) for v in parts.values()))
            for k in range(n):
                vals = [a.record.code, c.on_date, c.amount if k == 0 else None, uz.CATEGORY[c.category],
                        c.document, c.comment]
                for fw in FRAMEWORKS:
                    x = parts[fw][k] if k < len(parts[fw]) else None
                    if x is None:
                        vals += [uz.DECISION_SHORT["STANDART QO'LLANILMAYDI"] if not parts[fw] and k == 0 else None,
                                 None, None, None, None]
                    else:
                        vals += [uz.TREATMENT[x["treatment"]], x["amount"], rule_code(x["reason"]),
                                 reason_text(x["reason"]), rule_ref(x["reason"], fw)]
                r = _put(ws, r, vals, money_cols=(3, 8, 13), date_cols=(2,))
                for col in (7, 12):
                    v = ws.cell(r - 1, col).value
                    if v:
                        ws.cell(r - 1, col).font = GREEN if v == uz.TREATMENT["capitalize"] else RED
    _widths(ws, [16, 12, 16, 30, 14, 30, 16, 16, 8, 50, 16, 16, 16, 8, 50, 16])
    ws.auto_filter.ref = f"A{head}:P{max(r - 1, head)}"
    wb.save(path)


# ---------------------------------------------------------------- M10
def export_findings(analyses: list[Analysis], path, settings: AppSettings, signer, at: date):
    wb = Workbook()
    ws = wb.active
    ws.title = "Istisnolar"
    r = _title(ws, "Nazorat istisnolari — audit qo'mitasi uchun", settings, f"{uz.fdate(at)} holatiga",
               _demo(analyses))
    _header(ws, r, ["Daraja", "Kod", "Loyiha", "Qoida", "Standart", "Tavsif", "Sana"])
    r += 1
    order = {"high": 0, "medium": 1, "info": 2}
    items = sorted(((a.record, f) for a in analyses for f in a.findings),
                   key=lambda x: (order[x[1].level], int(x[1].code[1:]), x[0].code))
    for rec, f in items:
        r = _put(ws, r, [uz.LEVEL[f.level], f.code, rec.code, uz.CONTROL_TITLE[f.code],
                         uz.FW[f.framework] if f.framework else "", f.message, f.on], date_cols=(7,))
        if f.level == "high":
            ws.cell(r - 1, 1).font = RED
    r += 1
    ws.cell(r, 1, f"Jami istisnolar: {len(items)}").font = Font(bold=True)
    ws.cell(r + 2, 1, "Imzo:")
    ws.cell(r + 2, 2, f"{signer.name}, {uz.ROLE.get(signer.role, signer.role)}, {uz.fdate(clock.today())}")
    ws.cell(r + 3, 2, "______________________")
    _widths(ws, [10, 6, 16, 44, 10, 110, 12])
    wb.save(path)



# ---------------------------------------------------------------- M11
def export_audit_log(rows, path, settings: AppSettings):
    """Audit log; time shown in Tashkent (UTC+5), stored in UTC."""
    import json
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit jurnali"
    r = _title(ws, "Audit jurnali", settings, "Har bir o'zgarish: kim, qachon, nima edi → nima bo'ldi")
    _header(ws, r, [uz.t("col.when"), uz.t("col.who"), uz.t("col.role"), uz.t("col.action"), uz.t("col.object"),
                    uz.t("col.before"), uz.t("col.after")])
    r += 1
    for x in rows:
        obj = f"{x.entity} #{x.entity_id}" if x.entity_id is not None else x.entity
        r = _put(ws, r, [clock.fmt_tashkent(x.at), x.actor, uz.ROLE.get(x.role, x.role), x.action, obj,
                         json.dumps(x.before, ensure_ascii=False) if x.before is not None else "",
                         json.dumps(x.after, ensure_ascii=False) if x.after is not None else ""])
    _widths(ws, [21, 24, 26, 10, 20, 60, 60])
    wb.save(path)
