"""M9. Note to the financial statements (.docx): BHMS 11 (2024), 20-band and IAS 38.118 / 38.126.
Uzbek (Latin). Every number is taken from the engine (journal entries / schedules)."""
from __future__ import annotations

from datetime import date, datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .. import clock
from ..core.appsettings import AppSettings
from ..core.defaults import ACCOUNT_NAMES
from ..i18n import uz
from ..services import Analysis, movement, movement_total

FONT = "Times New Roman"


def _m(v) -> str:
    return uz.money(v or 0)


def _table(doc, headers, rows, widths=None, money_from: int = 1):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
    for row in rows:
        cells = t.add_row().cells
        bold = isinstance(row, tuple) and len(row) == 2 and row[1] == "bold"
        if bold:
            row = row[0]
        for j, v in enumerate(row):
            p = cells[j].paragraphs[0]
            r = p.add_run(str(v))
            r.bold = bold
            if j >= money_from and isinstance(v, str) and v.replace(" ", "").lstrip("-").isdigit():
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()
    return t


def _ready_by(a: Analysis, end: date) -> bool:
    """Asset available for use by the end of the period (and not derecognised before it)."""
    p = a.record
    return bool(p.available_for_use and p.available_for_use <= end and not (p.abandoned_on and p.abandoned_on <= end))


def _acc(code):
    return f"{code} «{ACCOUNT_NAMES[code]}»" if code in ACCOUNT_NAMES else code


def export_note(analyses: list[Analysis], path, settings: AppSettings, start: date, end: date):
    acc = settings.accounts
    inscope = [a for a in analyses if a.in_scope]
    outscope = [a for a in analyses if not a.in_scope]
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        hs = doc.styles[name]
        hs.font.name = FONT
        hs.font.color.rgb = RGBColor(0x1F, 0x2D, 0x3D)

    h = doc.add_heading("Moliyaviy hisobotga izoh: ilmiy-tadqiqot va tajriba-konstruktorlik ishlari (ITTKI) "
                        "xarajatlari", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"{settings.organization} · Hisobot davri: {uz.fdate(start)} — {uz.fdate(end)}")
    if any(a.record.is_demo for a in analyses):
        p = doc.add_paragraph()
        r = p.add_run(f"{uz.DEMO_MARK}. Raqamlar shartli, haqiqiy ma'lumotlar emas.")
        r.italic = True

    m_nas = movement_total(inscope, "NAS11", start, end, acc)
    m_ias = movement_total(inscope, "IAS38", start, end, acc)
    prod = acc.amortization_debit.startswith("2")       # 2010 / 2510: amortization goes into other assets' cost

    # ================================================================ BHMS 11
    doc.add_heading("1. BHMS 11 (2024) «Ilmiy-tadqiqot va tajriba-konstruktorlik ishlari bo'yicha xarajatlar», "
                    "20-band bo'yicha ochib berish", level=2)
    doc.add_heading("1.1. Hisob siyosati", level=3)
    for text in (
        "Tadqiqot bosqichidagi xarajatlar yuzaga kelgan davrda xarajat sifatida tan olinadi (BHMS 11, 10-band). "
        "Tadqiqot va ishlanma bosqichlarini ajratib bo'lmagan loyihalar bo'yicha barcha xarajatlar tadqiqot "
        "xarajati sifatida qaraladi (ehtiyotkorlik tamoyili).",
        "Ishlanma bosqichidagi xarajatlar BHMS 11 ning 12-bandidagi barcha mezonlar bajarilgan va dalil "
        "hujjatlari bilan tasdiqlangan sanadan boshlab kapitallashtiriladi. Mezonlar bajarilgunga qadar xarajatga "
        "o'tkazilgan summalar keyinchalik kapitallashtirilmaydi (11-band). Kapitallashtiriladigan summa kutilayotgan "
        "sof iqtisodiy naf bilan cheklanadi (13-band). Sotish xarajatlari aktiv qiymatiga kiritilmaydi (9-band).",
        f"Kapitallashtiriladigan xarajatlar {_acc(acc.capitalize)} hisobvarag'ida, foydalanishga tayyor aktiv "
        f"{_acc(acc.asset)} hisobvarag'ida, jamlangan amortizatsiya {_acc(acc.accumulated)} hisobvarag'ida hisobga "
        f"olinadi; xarajat sifatida tan olingan ITTKI xarajatlari — {_acc(acc.expense)}.",
        "Amortizatsiya to'g'ri chiziqli usulda, foydali xizmat muddati va 5 yildan (60 oy) kichigi davomida "
        f"hisoblanadi (BHMS 11, 15-16-band); boshlanishi: {uz.t('amort_start.' + settings.engine.amort_start).lower()}. "
        "Mezonlar keyinchalik bajarilmay qolsa ('Yo'q'), kapitallashtirilgan summa hisobdan chiqariladi (18-band).",
    ):
        doc.add_paragraph(text)

    doc.add_heading("1.2. Davr mobaynida xarajat sifatida tan olingan ITTKI xarajatlari", level=3)
    rows = []
    for a in inscope:
        mm = movement(a, "NAS11", start, end, acc)
        if mm["expensed_rnd"] or mm["expensed_other"]:
            rows.append([a.record.code, _m(mm["expensed_rnd"]), _m(mm["expensed_other"])])
    rows.append(([f"Jami", _m(m_nas["expensed_rnd"]), _m(m_nas["expensed_other"])], "bold"))
    _table(doc, ["Loyiha", "ITTKI xarajatlari, so'm", "ITTKI ga kirmaydigan (R1), so'm"], rows)
    doc.add_paragraph(f"Davr mobaynida xarajat sifatida tan olingan ITTKI xarajatlari: {_m(m_nas['expensed_rnd'])} "
                      f"so'm. Loyihalarga yozilgan, ammo aktiv qiymatiga kiritilmaydigan toifalar (sotish, umumiy "
                      f"ma'muriy, o'qitish va boshqalar): {_m(m_nas['expensed_other'])} so'm.")

    doc.add_heading("1.3. Amortizatsiya usullari, muddatlari va normalari", level=3)
    rows = []
    for a in inscope:
        n = a.res["NAS11"]["amort_months"]
        if n and _ready_by(a, end):
            rows.append([a.record.code, "to'g'ri chiziqli", str(a.record.useful_life_months), str(n),
                         f"{1200 / n:.1f}%".replace(".", ",")])
    if rows:
        _table(doc, ["Loyiha", "Usul", "Foydali xizmat muddati, oy", "BHMS 11 muddati, oy", "Yillik norma"], rows,
               money_from=99)
    else:
        doc.add_paragraph("Davr oxiriga foydalanishga tayyor ITTKI natijalari yo'q.")

    doc.add_heading("1.4. Kapitallashtirilgan ITTKI xarajatlari harakati", level=3)
    written_off = m_nas["wip_writeoff"] + m_nas["impairment"] + m_nas["disposal_residual"]
    amort_exp = 0 if prod else m_nas["amortization"]
    amort_assets = m_nas["amortization"] if prod else 0
    _table(doc, ["Ko'rsatkich", "so'm"], [
        ([f"Davr boshiga qoldiq ({uz.fdate(start)})", _m(m_nas["carrying_open"])], "bold"),
        ["Kapitallashtirildi", _m(m_nas["capitalized"])],
        ["Amortizatsiya hisoblandi (xarajatga)", _m(amort_exp)],
        ["Boshqa aktivlar qiymatiga o'tkazildi (amortizatsiya, 17-band)", _m(amort_assets)],
        ["Hisobdan chiqarildi (18-band, loyiha to'xtatilishi, qadrsizlanish)", _m(written_off)],
        ["Tiklandi (19-band)", _m(m_nas["restored"])],
        ([f"Davr oxiriga qoldiq ({uz.fdate(end)})", _m(m_nas["carrying_close"])], "bold"),
    ])
    doc.add_paragraph(f"Shu jumladan davr oxiriga: yaratilayotgan (foydalanishga tayyor bo'lmagan) — "
                      f"{_m(m_nas['wip_close'])} so'm; foydalanishga tayyor — boshlang'ich qiymati "
                      f"{_m(m_nas['gross_close'])} so'm, jamlangan amortizatsiya {_m(m_nas['accum_close'])} so'm. "
                      f"Davr mobaynida foydalanishga tayyor aktivlarga o'tkazildi: {_m(m_nas['transfer'])} so'm.")

    # ================================================================ IAS 38
    doc.add_heading("2. IAS 38 «Nomoddiy aktivlar» bo'yicha ochib berish (38.118, 38.126)", level=2)
    doc.add_heading("2.1. Ichki yaratilgan nomoddiy aktivlar sinfi (IAS 38.118)", level=3)
    rows = []
    for a in inscope:
        n = a.res["IAS38"]["amort_months"]
        if n and _ready_by(a, end):
            rows.append([a.record.code, "chekli", str(n), "to'g'ri chiziqli", f"{1200 / n:.1f}%".replace(".", ",")])
    if rows:
        _table(doc, ["Loyiha", "Foydali xizmat muddati", "Muddat, oy", "Usul (38.118(a)-(b))", "Yillik norma"],
               rows, money_from=99)
    line = ACCOUNT_NAMES.get(acc.amortization_debit, acc.amortization_debit)
    doc.add_paragraph(f"Amortizatsiya foyda va zarar to'g'risidagi hisobotning «{line}» moddasiga kiritiladi "
                      f"(hisobvaraq {acc.amortization_debit}) (38.118(d)).")
    _table(doc, ["Ko'rsatkich (38.118(c), (e))", "so'm"], [
        ([f"Boshlang'ich qiymat, davr boshiga", _m(m_ias["gross_open"])], "bold"),
        ["Jamlangan amortizatsiya va qadrsizlanish, davr boshiga", _m(m_ias["accum_open"])],
        ([f"Balans qiymati, davr boshiga", _m(m_ias["gross_open"] - m_ias["accum_open"])], "bold"),
        ["Qo'shildi: ichki ishlanmadan (38.118(e)(i))", _m(m_ias["transfer"])],
        ["Chiqib ketdi (38.118(e)(ii))", _m(m_ias["disposal_residual"])],
        ["Qadrsizlanishdan zarar (38.118(e)(iv))", _m(m_ias["impairment"])],
        ["Amortizatsiya (38.118(e)(vi))", _m(m_ias["amortization"])],
        ([f"Boshlang'ich qiymat, davr oxiriga", _m(m_ias["gross_close"])], "bold"),
        ["Jamlangan amortizatsiya va qadrsizlanish, davr oxiriga", _m(m_ias["accum_close"])],
        ([f"Balans qiymati, davr oxiriga", _m(m_ias["gross_close"] - m_ias["accum_close"])], "bold"),
    ])
    doc.add_heading("2.2. Yaratilayotgan (foydalanishga tayyor bo'lmagan) nomoddiy aktivlar", level=3)
    _table(doc, ["Ko'rsatkich", "so'm"], [
        (["Davr boshiga", _m(m_ias["wip_open"])], "bold"),
        ["Kapitallashtirildi (IAS 38.57)", _m(m_ias["capitalized"])],
        ["Qadrsizlanish / hisobdan chiqarish (IAS 36, 38.112)", _m(m_ias["wip_writeoff"])],
        ["Foydalanishga tayyor aktivlarga o'tkazildi", _m(m_ias["transfer"])],
        (["Davr oxiriga", _m(m_ias["wip_close"])], "bold"),
    ])
    doc.add_paragraph("Foydalanishga tayyor bo'lmagan nomoddiy aktivlar har yili qadrsizlanishga tekshiriladi "
                      "(IAS 36.10(a)).")
    doc.add_heading("2.3. Xarajat sifatida tan olingan tadqiqot va ishlanma xarajatlari (IAS 38.126)", level=3)
    doc.add_paragraph(f"Davr mobaynida xarajat sifatida tan olingan tadqiqot va ishlanma xarajatlarining umumiy "
                      f"summasi: {_m(m_ias['expensed_rnd'])} so'm.")

    # ================================================================ differences
    doc.add_heading("3. BHMS 11 va IAS 38 o'rtasidagi farqlar", level=2)
    rows = []
    for a in inscope:
        rows.append([a.record.code, _m(a.carrying["IAS38"]), _m(a.carrying["NAS11"]),
                     _m(a.carrying["IAS38"] - a.carrying["NAS11"])])
    _table(doc, ["Loyiha", f"IAS 38, {uz.fdate(end)}", f"BHMS 11, {uz.fdate(end)}", "Farq"], rows)
    for a in inscope:
        if a.reasons and a.reasons != [uz.REC_REASON["none"]]:
            doc.add_paragraph(f"{a.record.code}:")
            for r in a.reasons:
                doc.add_paragraph(r, style="List Bullet")
    if outscope:
        doc.add_paragraph("Standart qo'llanilmaydigan loyihalar (BHMS 11, 4-5-band; IAS 38.2): "
                          + ", ".join(a.record.code for a in outscope) + ".")

    foot = doc.add_paragraph()
    r = foot.add_run(f"Izoh {uz.APP_NAME} dasturida {clock.now_tashkent()} (Toshkent, UTC+5) da avtomatik tuzildi; raqamlar "
                     f"qarorlar mexanizmi hisob-kitoblaridan olingan.")
    r.italic = True
    r.font.size = Pt(9)
    doc.save(path)
